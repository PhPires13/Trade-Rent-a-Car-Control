"""Etapa 17 (retomada) — arquivo da contabilidade com a lista de pagamentos do DAS
e franquia sem data do evento contando pela data do sinistro."""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.financeiro.models import AplicacaoRecebimento, Cobranca, Recebimento
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente
from apps.relatorios import services
from apps.sinistros.models import Sinistro


@pytest.fixture
def cliente(db):
    return Cliente.objects.create(nome="Arlen", cpf_cnpj="111.222.333-44")


@pytest.fixture
def veiculo(db):
    return Veiculo.objects.create(placa="QXQ7R18", marca_modelo="Gol")


def _pagar(cliente, valor, data, descricao):
    cobranca = Cobranca.objects.create(
        cliente=cliente, descricao=descricao, valor=valor, vencimento=data
    )
    recebimento = Recebimento.objects.create(cliente=cliente, data=data, valor=valor)
    AplicacaoRecebimento.objects.create(recebimento=recebimento, cobranca=cobranca, valor=valor)
    return cobranca


def test_arquivo_da_contabilidade_traz_os_pagamentos_do_das(usuario_logado, cliente):
    _pagar(cliente, Decimal("650.00"), date(2026, 9, 8), "Aluguel semana 1")
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-09", "exportar": "tudo"})
    assert resposta.status_code == 200
    wb = load_workbook(io.BytesIO(resposta.content))
    assert "Pagamentos do DAS" in wb.sheetnames
    linhas = list(wb["Pagamentos do DAS"].iter_rows(values_only=True))
    assert linhas[0][0] == "Data" and linhas[0][5] == "Valor (R$)"
    pagamento = linhas[1]
    assert pagamento[0] == "08/09/2026"
    assert pagamento[1] == "Arlen"
    assert pagamento[3] == "Aluguel semana 1"
    assert pagamento[5] == 650.0
    # total da aba bate com o total da aba DAS (locação + diversos)
    assert linhas[-2][3].startswith("Total") and linhas[-2][5] == 650.0


def test_pagamento_de_outro_mes_nao_entra_na_aba(usuario_logado, cliente):
    _pagar(cliente, Decimal("650.00"), date(2026, 8, 30), "Aluguel agosto")
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-09", "exportar": "tudo"})
    linhas = list(
        load_workbook(io.BytesIO(resposta.content))["Pagamentos do DAS"].iter_rows(values_only=True)
    )
    assert len(linhas) == 3  # cabeçalho + as duas linhas de total (zeradas)
    assert linhas[-2][5] == 0


def test_franquia_sem_data_do_evento_entra_no_mes_do_sinistro(veiculo):
    Sinistro.objects.create(
        veiculo=veiculo,
        data=date(2026, 4, 12),
        envolvido="terceiro",
        acionou_protecao=True,
        franquia_valor=Decimal("800.00"),
    )
    assert services.despesas_do_mes(2026, 4)["total_franquias"] == Decimal("800.00")
    assert services.despesas_do_mes(2026, 5)["total_franquias"] == Decimal("0")
