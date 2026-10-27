"""Qualidade final dos Relatórios: franquia sem data do evento, proteção de carro
sem data de aquisição, CSV em português, planilha da frota legível, card de
receitas, quadros "de hoje", escolha de mês e o arquivo único da contabilidade."""

import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.financeiro.models import Cobranca
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente
from apps.relatorios import services
from apps.sinistros.models import Sinistro


@pytest.fixture
def veiculo(db):
    return Veiculo.objects.create(
        placa="QXQ6C10",
        marca_modelo="Gol",
        mensalidade_protecao=Decimal("304.00"),
        data_aquisicao=date(2025, 1, 1),
    )


@pytest.fixture
def cliente(db):
    return Cliente.objects.create(nome="Arlen", cpf_cnpj="111.222.333-44")


def _csv(resposta):
    return list(csv.reader(io.StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"))


# --- franquia sem data do evento -------------------------------------------------


def _franquia_sem_data_evento(veiculo):
    return Sinistro.objects.create(
        veiculo=veiculo,
        data=date(2026, 3, 5),
        envolvido="terceiro",
        acionou_protecao=True,
        franquia_valor=Decimal("1500.00"),
    )


def test_franquia_sem_data_do_evento_conta_pela_data_do_sinistro(veiculo):
    _franquia_sem_data_evento(veiculo)
    despesas = services.despesas_do_mes(2026, 3)
    assert despesas["total_franquias"] == Decimal("1500.00")
    serie = services.serie_mensal(2026, 3, quantidade=1)[0]
    assert serie["despesa"] == float(Decimal("1500.00") + despesas["total_protecao"])


def test_franquia_com_data_do_evento_conta_pela_data_do_evento(veiculo):
    Sinistro.objects.create(
        veiculo=veiculo,
        data=date(2026, 2, 25),
        envolvido="terceiro",
        acionou_protecao=True,
        data_evento=date(2026, 3, 2),
        franquia_valor=Decimal("900.00"),
    )
    assert services.despesas_do_mes(2026, 2)["total_franquias"] == Decimal("0")
    assert services.despesas_do_mes(2026, 3)["total_franquias"] == Decimal("900.00")


def test_franquia_sem_data_do_evento_sai_na_planilha_de_despesas(usuario_logado, veiculo):
    _franquia_sem_data_evento(veiculo)
    resposta = usuario_logado.get(
        "/relatorios/", {"mes": "2026-03", "exportar": "despesas", "formato": "xlsx"}
    )
    ws = load_workbook(io.BytesIO(resposta.content)).active
    linhas = [list(linha) for linha in ws.iter_rows(min_row=2, values_only=True)]
    franquia = next(linha for linha in linhas if linha[2] == "Franquia de evento")
    assert franquia[0] == "05/03/2026"
    assert franquia[4] == 1500.0


# --- proteção de carro sem data de aquisição ------------------------------------


def test_carro_sem_data_de_aquisicao_nao_lanca_protecao(db):
    Veiculo.objects.create(
        placa="SEM1D00", marca_modelo="Onix", mensalidade_protecao=Decimal("304.00")
    )
    assert services.despesas_do_mes(2026, 3)["total_protecao"] == Decimal("0")
    assert services.serie_mensal(2026, 3, quantidade=2)[-1]["despesa"] == 0.0


def test_carro_com_data_de_aquisicao_lanca_protecao_so_depois_dela(db):
    Veiculo.objects.create(
        placa="COM1D00",
        marca_modelo="Onix",
        mensalidade_protecao=Decimal("304.00"),
        data_aquisicao=date(2026, 3, 10),
    )
    assert services.despesas_do_mes(2026, 2)["total_protecao"] == Decimal("0")
    assert services.despesas_do_mes(2026, 3)["total_protecao"] == Decimal("304.00")
    serie = services.serie_mensal(2026, 3, quantidade=2)
    assert [p["despesa"] for p in serie] == [0.0, 304.0]


# --- CSV em português e só o botão do Excel na tela ------------------------------


def test_csv_tem_bom_ponto_e_virgula_e_virgula_decimal(usuario_logado, veiculo):
    resposta = usuario_logado.get(
        "/relatorios/", {"mes": "2026-03", "exportar": "despesas", "formato": "csv"}
    )
    assert resposta.content.startswith("﻿".encode())
    linhas = _csv(resposta)
    assert linhas[0] == ["Data", "Veículo", "Origem", "Descrição", "Valor (R$)"]
    assert linhas[-1] == ["", "", "TOTAL", "", "304,00"]


def test_tela_so_tem_o_botao_do_excel(usuario_logado, db):
    html = usuario_logado.get("/relatorios/?mes=2026-03").content.decode()
    assert "formato=csv" not in html
    assert "⬇ Baixar planilha" in html


# --- planilha da frota legível ---------------------------------------------------


def test_planilha_da_frota_tem_percentual_de_0_a_100_e_recomendacao_legivel(
    usuario_logado, veiculo
):
    Veiculo.objects.create(placa="SEM1V00", marca_modelo="Onix")  # sem investimento
    resposta = usuario_logado.get("/relatorios/", {"exportar": "frota", "formato": "xlsx"})
    ws = load_workbook(io.BytesIO(resposta.content)).active
    cabecalho = [c.value for c in ws[1]]
    assert "Já voltou (%)" in cabecalho
    i_perc = cabecalho.index("Já voltou (%)")
    i_rec = cabecalho.index("Recomendação")
    linhas = {linha[0]: linha for linha in ws.iter_rows(min_row=2, values_only=True)}
    assert linhas["SEM1V00"][i_perc] is None  # sem investimento: vazio, não 0
    assert linhas["SEM1V00"][i_rec] in {"Manter", "Observar", "Preparar venda", "Vender"}
    assert linhas["QXQ6C10"][i_rec] in {"Manter", "Observar", "Preparar venda", "Vender"}


def test_planilha_da_frota_percentual_em_escala_de_cem(usuario_logado, veiculo, monkeypatch):
    from apps.relatorios import views

    fichas, media = services.resumo_da_frota()
    ficha = next(f for f in fichas if f.veiculo.placa == "QXQ6C10")
    monkeypatch.setattr(
        type(ficha), "percentual_recuperado", property(lambda self: Decimal("0.756"))
    )
    monkeypatch.setattr(views.services, "resumo_da_frota", lambda: ([ficha], media))
    _, colunas, linhas = views._planilha_frota(2026, 10)
    assert linhas[0][colunas.index("Já voltou (%)")] == 75.6


# --- card de receitas e link para o DAS ------------------------------------------


def test_card_de_receitas_explica_a_diferenca_com_o_caixa_do_das(usuario_logado, db):
    html = usuario_logado.get("/relatorios/?mes=2026-03").content.decode()
    assert '<p class="text-sm text-green-800">Receitas do mês</p>' in html
    assert ">Entrou<" not in html
    assert "/financeiro/das/?mes=2026-03" in html
    assert "Total que entrou no caixa" in html


# --- quadros de hoje -------------------------------------------------------------


def test_recebiveis_e_frota_dizem_que_sao_de_hoje(usuario_logado, db):
    hoje = timezone.localdate()
    html = usuario_logado.get("/relatorios/?mes=2026-03").content.decode()
    assert f"(hoje, {hoje:%d/%m/%Y})" in html
    assert "Resumo da frota" in html and "(hoje)" in html


@pytest.mark.parametrize("tipo", ["recebiveis", "frota"])
def test_arquivo_de_hoje_leva_a_data_de_hoje(usuario_logado, db, tipo):
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-03", "exportar": tipo})
    assert f"{tipo}-{timezone.localdate().isoformat()}.xlsx" in resposta["Content-Disposition"]


def test_arquivo_do_mes_leva_o_mes(usuario_logado, db):
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-03", "exportar": "receitas"})
    assert "receitas-2026-03.xlsx" in resposta["Content-Disposition"]


def _judicial(cliente, valor, status=Cobranca.Status.JUDICIAL, descricao="Dívida"):
    Cobranca.objects.create(
        cliente=cliente,
        origem="aluguel",
        descricao=descricao,
        valor=valor,
        vencimento=date(2026, 1, 1),
        status=status,
    )


def test_cliente_todo_na_justica_mostra_so_o_selo(usuario_logado, cliente):
    _judicial(cliente, Decimal("2000.00"))
    html = usuario_logado.get("/relatorios/").content.decode()
    assert ">na Justiça<" in html
    assert "judicial R$" not in html
    assert "dos quais" not in html


def test_cliente_parte_na_justica_mostra_quanto(usuario_logado, cliente):
    _judicial(cliente, Decimal("2000.00"))
    _judicial(cliente, Decimal("300.00"), status=Cobranca.Status.ATRASADO, descricao="Semana")
    html = usuario_logado.get("/relatorios/").content.decode()
    assert "dos quais R$ 2.000,00 na Justiça" in html


# --- escolha de mês --------------------------------------------------------------


def test_caixinha_de_mes_vai_direto_ao_mes_escolhido(usuario_logado, db):
    html = usuario_logado.get("/relatorios/?mes=2025-04").content.decode()
    assert "data-escolher-mes" in html
    assert '<option value="2025-04" selected>Abril de 2025</option>' in html
    hoje = timezone.localdate()
    assert f'value="{hoje.year}-{hoje.month:02d}"' in html


@pytest.mark.parametrize("mes", ["2026-13", "banana", "0001-01"])
def test_mes_invalido_avisa(usuario_logado, db, mes):
    resposta = usuario_logado.get(f"/relatorios/?mes={mes}")
    hoje = timezone.localdate()
    assert (resposta.context["ano"], resposta.context["mes"]) == (hoje.year, hoje.month)
    assert "Mês inválido — mostrando o mês atual" in resposta.content.decode()


def test_sem_mes_na_url_nao_avisa(usuario_logado, db):
    assert "Mês inválido" not in usuario_logado.get("/relatorios/").content.decode()


def test_mes_valido_nao_avisa(usuario_logado, db):
    assert "Mês inválido" not in usuario_logado.get("/relatorios/?mes=2026-03").content.decode()


# --- arquivo único para a contabilidade ------------------------------------------


def test_baixar_o_mes_para_a_contabilidade_gera_um_arquivo_com_todas_as_abas(
    usuario_logado, veiculo
):
    html = usuario_logado.get("/relatorios/?mes=2026-03").content.decode()
    assert "Baixar o mês para a contabilidade" in html
    assert "?mes=2026-03&amp;exportar=tudo" in html or "?mes=2026-03&exportar=tudo" in html
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-03", "exportar": "tudo"})
    assert resposta.status_code == 200
    assert "contabilidade-2026-03.xlsx" in resposta["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resposta.content))
    nomes = wb.sheetnames
    assert nomes[:3] == ["Receitas", "Despesas", "DAS"]
    assert nomes[3].startswith("Recebíveis hoje")
    assert nomes[4].startswith("Frota hoje")
    despesas = [linha for linha in wb["Despesas"].iter_rows(values_only=True)]
    assert despesas[-1][2] == "TOTAL" and despesas[-1][4] == 304.0
    das = [linha[0] for linha in wb["DAS"].iter_rows(values_only=True)]
    assert "Total que entrou no caixa" in das


# --- revisão: gráfico soma franquias com e sem data do evento, mês a mês ---------


def test_grafico_soma_franquias_com_e_sem_data_do_evento_no_mes_certo(veiculo):
    _franquia_sem_data_evento(veiculo)  # 1500 em março (data do sinistro)
    Sinistro.objects.create(  # 900 em março (data do evento), sinistro em fevereiro
        veiculo=veiculo,
        data=date(2026, 2, 25),
        envolvido="terceiro",
        acionou_protecao=True,
        data_evento=date(2026, 3, 2),
        franquia_valor=Decimal("900.00"),
    )
    Sinistro.objects.create(  # 200 em janeiro
        veiculo=veiculo,
        data=date(2026, 1, 9),
        envolvido="terceiro",
        acionou_protecao=True,
        franquia_valor=Decimal("200.00"),
    )
    serie = services.serie_mensal(2026, 3, quantidade=3)
    protecao = 304.0
    assert [p["despesa"] for p in serie] == [200 + protecao, protecao, 2400 + protecao]
    assert services.despesas_do_mes(2026, 3)["total_franquias"] == Decimal("2400.00")


def test_mes_atual_usa_o_dia_de_hoje_do_fuso_de_bh(usuario_logado, db):
    from django.utils import timezone

    hoje = timezone.localdate()
    html = usuario_logado.get("/relatorios/").content.decode()
    assert f"(hoje, {hoje:%d/%m/%Y})" in html
    assert "Mês inválido" not in html
