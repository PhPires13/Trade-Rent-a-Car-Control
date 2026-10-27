"""Etapa 17 (extra) — relatórios: título da tela, ?mes= com ano gigante e a
mesma planilha do DAS em Financeiro › DAS e em Relatórios."""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.alocacoes.models import Alocacao
from apps.financeiro.models import Caucao, MovimentacaoCaucao
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente


def test_titulo_diz_que_e_a_tela_de_relatorios(usuario_logado):
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-10"})
    assert "Relatórios — Outubro de 2026</h1>" in resposta.content.decode()


@pytest.mark.parametrize("url", ["/relatorios/", "/financeiro/das/", "/relatorios/?exportar=tudo"])
def test_mes_com_ano_gigante_nao_quebra(usuario_logado, url):
    separador = "&" if "?" in url else "?"
    resposta = usuario_logado.get(f"{url}{separador}mes=99999999999999999999-01")
    assert resposta.status_code == 200


@pytest.fixture
def caucao_de_setembro(db):
    veiculo = Veiculo.objects.create(placa="QXQ7R28", marca_modelo="Gol")
    cliente = Cliente.objects.create(nome="Arlen", cpf_cnpj="111.222.333-44")
    alocacao = Alocacao.objects.create(
        veiculo=veiculo,
        cliente=cliente,
        data_inicio=date(2026, 8, 1),
        valor_semanal=Decimal("650.00"),
        km_entrega=10_000,
    )
    caucao = Caucao.objects.create(alocacao=alocacao)
    MovimentacaoCaucao.objects.create(
        caucao=caucao,
        tipo=MovimentacaoCaucao.Tipo.RECEBIMENTO,
        data=date(2026, 9, 1),
        valor=Decimal("500.00"),
    )


def _abas(resposta):
    assert resposta.status_code == 200
    livro = load_workbook(io.BytesIO(resposta.content))
    return [[list(linha) for linha in aba.iter_rows(values_only=True)] for aba in livro]


def test_planilha_do_das_igual_em_financeiro_e_relatorios(usuario_logado, caucao_de_setembro):
    financeiro = _abas(usuario_logado.get("/financeiro/das/?mes=2026-09&exportar=xlsx"))
    relatorios = _abas(usuario_logado.get("/relatorios/?mes=2026-09&exportar=tudo"))
    # Relatórios: 3ª aba = DAS, última = Pagamentos do DAS
    assert financeiro[0] == relatorios[2]
    assert financeiro[1] == relatorios[-1]
    pagamentos = financeiro[1]
    assert any(linha[3] == "Entrou no mês, mas fora do DAS:" for linha in pagamentos)
    caucao = [linha for linha in pagamentos if linha[3] and "Caução recebida" in linha[3]]
    assert caucao and caucao[0][5] == 500 and caucao[0][6] == "Não"
