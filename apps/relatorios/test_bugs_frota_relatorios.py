"""Correções da frente frota-relatórios nos relatórios: C10, C12, C23, D06, D07, D14."""

import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.frota.models import PagamentoIpva, Veiculo
from apps.manutencao.models import Manutencao
from apps.multas.models import Multa
from apps.relatorios import services
from apps.sinistros.models import Sinistro


@pytest.fixture
def veiculo(db):
    return Veiculo.objects.create(
        placa="QXQ6C10",
        marca_modelo="Gol",
        data_aquisicao=date(2025, 1, 1),
        valor_compra=Decimal("40000.00"),
    )


def _despesa_da_serie(ano, mes):
    return services.serie_mensal(ano, mes, quantidade=1)[0]["despesa"]


# --- D06: IPVA do ano anterior não some ao lançar o ano novo ------------------------


def test_d06_lancar_ipva_do_ano_novo_nao_muda_o_mes_ja_enviado(veiculo):
    PagamentoIpva.objects.create(
        veiculo=veiculo, ano=2026, valor=Decimal("1200.00"), pago_em=date(2026, 1, 15)
    )
    assert services.despesas_do_mes(2026, 1)["total_ipva"] == Decimal("1200.00")
    PagamentoIpva.objects.create(
        veiculo=veiculo, ano=2027, valor=Decimal("1260.00"), pago_em=date(2027, 2, 10)
    )
    assert services.despesas_do_mes(2026, 1)["total_ipva"] == Decimal("1200.00")
    assert services.despesas_do_mes(2027, 2)["total_ipva"] == Decimal("1260.00")
    assert _despesa_da_serie(2026, 1) == 1200.0
    assert _despesa_da_serie(2027, 2) == 1260.0


def test_d06_ipva_parcelado_entra_no_mes_de_cada_parcela(veiculo, usuario_logado):
    for parcela, mes in ((1, 1), (2, 2), (3, 3)):
        PagamentoIpva.objects.create(
            veiculo=veiculo,
            ano=2026,
            parcela=parcela,
            valor=Decimal("400.00"),
            pago_em=date(2026, mes, 10),
        )
    assert services.despesas_do_mes(2026, 2)["total_ipva"] == Decimal("400.00")
    resposta = usuario_logado.get(
        "/relatorios/", {"mes": "2026-02", "exportar": "despesas", "formato": "csv"}
    )
    linhas = list(csv.reader(io.StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"))
    assert ["10/02/2026", "QXQ6C10", "IPVA", "IPVA 2026 — parcela 2", "400,00"] in linhas


# --- D07: cota da proteção só no sinistro -------------------------------------------


def test_d07_cota_da_protecao_nao_conta_em_dobro_no_mes(veiculo):
    sinistro = Sinistro.objects.create(
        veiculo=veiculo,
        data=date(2026, 3, 5),
        envolvido="terceiro",
        acionou_protecao=True,
        data_evento=date(2026, 3, 5),
        franquia_valor=Decimal("1500.00"),
    )
    Manutencao.objects.create(
        veiculo=veiculo,
        tipo="corretiva",
        data=date(2026, 3, 20),
        descricao="Funilaria do sinistro",
        origem_custo=Manutencao.OrigemCusto.PROTECAO,
        custo_real=Decimal("1500.00"),
        sinistro=sinistro,
    )
    Manutencao.objects.create(
        veiculo=veiculo,
        tipo="corretiva",
        data=date(2026, 3, 21),
        descricao="Farol (por fora)",
        custo_real=Decimal("300.00"),
        sinistro=sinistro,
    )
    despesas = services.despesas_do_mes(2026, 3)
    assert despesas["total_franquias"] == Decimal("1500.00")
    assert despesas["total_manutencao"] == Decimal("300.00")
    assert despesas["total_geral"] == Decimal("1800.00")
    assert _despesa_da_serie(2026, 3) == 1800.0


def test_d07_conserto_pela_protecao_sem_franquia_no_sinistro_continua_contando(veiculo):
    Manutencao.objects.create(
        veiculo=veiculo,
        tipo="corretiva",
        data=date(2026, 3, 20),
        descricao="Evento sem sinistro cadastrado",
        origem_custo=Manutencao.OrigemCusto.PROTECAO,
        custo_real=Decimal("900.00"),
    )
    assert services.despesas_do_mes(2026, 3)["total_manutencao"] == Decimal("900.00")


# --- C10: multa anulada não é despesa -----------------------------------------------


def test_c10_multa_anulada_sai_das_despesas_do_mes_e_da_serie(veiculo):
    anulada = Multa.objects.create(
        veiculo=veiculo,
        data_infracao=date(2026, 5, 3),
        valor=Decimal("195.23"),
        responsavel=Multa.Responsavel.EMPRESA,
    )
    assert services.despesas_do_mes(2026, 5)["total_multas_empresa"] == Decimal("195.23")
    anulada.resultado = Multa.Resultado.ADVERTENCIA
    anulada.save()
    despesas = services.despesas_do_mes(2026, 5)
    assert despesas["total_multas_empresa"] == Decimal("0")
    assert despesas["multas_empresa"] == []
    assert _despesa_da_serie(2026, 5) == 0.0


# --- D14: relatório da contabilidade só com carros de locação -----------------------


def test_d14_carro_pessoal_fica_fora_do_relatorio(db):
    pessoal = Veiculo.objects.create(
        placa="HBB2O20",
        marca_modelo="HB20",
        uso=Veiculo.Uso.FORA_LOCACAO,
        data_aquisicao=date(2026, 6, 2),
        custos_entrada=Decimal("300.00"),
    )
    Manutencao.objects.create(
        veiculo=pessoal,
        tipo="corretiva",
        data=date(2026, 6, 10),
        descricao="Revisão",
        custo_real=Decimal("800.00"),
    )
    PagamentoIpva.objects.create(
        veiculo=pessoal, ano=2026, valor=Decimal("900.00"), pago_em=date(2026, 6, 15)
    )
    Multa.objects.create(
        veiculo=pessoal,
        data_infracao=date(2026, 6, 20),
        valor=Decimal("130.16"),
        responsavel=Multa.Responsavel.EMPRESA,
    )
    despesas = services.despesas_do_mes(2026, 6)
    assert despesas["total_geral"] == Decimal("0")
    assert _despesa_da_serie(2026, 6) == 0.0
    # continua no acompanhamento do próprio carro
    from apps.frota.desmobilizacao import montar_ficha

    ficha = montar_ficha(pessoal, hoje=date(2026, 6, 30))
    assert ficha.despesa_manutencao == Decimal("800.00")
    assert ficha.despesa_ipva == Decimal("900.00")


def test_d14_venda_de_carro_pessoal_nao_e_receita_da_locadora(db):
    Veiculo.objects.create(
        placa="HBB2O20",
        marca_modelo="HB20",
        uso=Veiculo.Uso.FORA_LOCACAO,
        data_venda=date(2026, 6, 10),
        valor_venda=Decimal("35000.00"),
        status=Veiculo.Status.VENDIDO,
    )
    assert services.receitas_do_mes(2026, 6)["total_vendas"] == Decimal("0")
    assert services.serie_mensal(2026, 6, quantidade=1)[0]["receita"] == 0.0


# --- C12: export da frota com o IPVA ------------------------------------------------


def test_c12_export_da_frota_inclui_o_ipva_nas_despesas(usuario_logado, veiculo):
    PagamentoIpva.objects.create(
        veiculo=veiculo, ano=2026, valor=Decimal("1200.00"), pago_em=date(2026, 1, 15)
    )
    resposta = usuario_logado.get("/relatorios/", {"exportar": "frota", "formato": "csv"})
    linhas = list(csv.reader(io.StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"))
    linha = next(linha for linha in linhas if linha[0] == "QXQ6C10")
    assert linha[4] == "1200,00"  # CSV em português: vírgula nos centavos


# --- C23: texto que parece fórmula ---------------------------------------------------


def _manutencao_com_formula(veiculo):
    Manutencao.objects.create(
        veiculo=veiculo,
        tipo="corretiva",
        data=date(2026, 7, 10),
        descricao='=HIPERLINK("http://x";"ver nota")',
        custo_real=Decimal("100.00"),
    )


def test_c23_xlsx_grava_descricao_com_igual_como_texto(usuario_logado, veiculo):
    _manutencao_com_formula(veiculo)
    resposta = usuario_logado.get(
        "/relatorios/", {"mes": "2026-07", "exportar": "despesas", "formato": "xlsx"}
    )
    ws = load_workbook(io.BytesIO(resposta.content)).active
    celula = ws.cell(row=2, column=4)
    assert celula.data_type == "s"
    assert celula.value == '=HIPERLINK("http://x";"ver nota")'


def test_c23_csv_prefixa_apostrofo_em_texto_que_parece_formula(usuario_logado, veiculo):
    _manutencao_com_formula(veiculo)
    resposta = usuario_logado.get(
        "/relatorios/", {"mes": "2026-07", "exportar": "despesas", "formato": "csv"}
    )
    linhas = list(csv.reader(io.StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"))
    assert linhas[1][3] == '\'=HIPERLINK("http://x";"ver nota")'
    assert linhas[1][4] == "100,00"  # número continua número (vírgula decimal)
