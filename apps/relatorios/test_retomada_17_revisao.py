"""Etapa 17 (retomada) — revisão: a aba "Pagamentos do DAS" do arquivo da
contabilidade fecha com a aba "DAS" (inclusive desconto na caução) e não deixa
descrição digitada virar fórmula."""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.alocacoes.models import Alocacao
from apps.financeiro.models import (
    AplicacaoRecebimento,
    Caucao,
    Cobranca,
    MovimentacaoCaucao,
    Recebimento,
)
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente


@pytest.fixture
def alocacao(db):
    veiculo = Veiculo.objects.create(placa="QXQ7R28", marca_modelo="Gol")
    cliente = Cliente.objects.create(nome="Arlen", cpf_cnpj="111.222.333-44")
    return Alocacao.objects.create(
        veiculo=veiculo,
        cliente=cliente,
        data_inicio=date(2026, 8, 1),
        valor_semanal=Decimal("650.00"),
        km_entrega=10_000,
    )


def _abas(usuario_logado):
    resposta = usuario_logado.get("/relatorios/", {"mes": "2026-09", "exportar": "tudo"})
    assert resposta.status_code == 200
    return load_workbook(io.BytesIO(resposta.content))


def test_pagamentos_do_das_fecham_com_a_aba_das(usuario_logado, alocacao):
    cliente = alocacao.cliente
    aluguel = Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel semana 1",
        valor=Decimal("650.00"),
        vencimento=date(2026, 9, 8),
    )
    recebimento = Recebimento.objects.create(
        cliente=cliente, data=date(2026, 9, 8), valor=Decimal("650.00")
    )
    AplicacaoRecebimento.objects.create(
        recebimento=recebimento, cobranca=aluguel, valor=Decimal("650.00")
    )
    # ND paga com desconto na caução: entra na lista, mas fora da base do DAS
    nd = Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.NOTA_DEBITO,
        descricao='=HIPERLINK("x")',
        valor=Decimal("120.00"),
        vencimento=date(2026, 9, 15),
    )
    caucao = Caucao.objects.create(alocacao=alocacao)
    MovimentacaoCaucao.objects.create(
        caucao=caucao,
        tipo=MovimentacaoCaucao.Tipo.RECEBIMENTO,
        valor=Decimal("500.00"),
        data=date(2026, 8, 1),
    )
    MovimentacaoCaucao.objects.create(
        caucao=caucao,
        tipo=MovimentacaoCaucao.Tipo.DESCONTO,
        valor=Decimal("120.00"),
        data=date(2026, 9, 15),
        cobranca=nd,
    )

    wb = _abas(usuario_logado)
    linhas = list(wb["Pagamentos do DAS"].iter_rows(values_only=True))
    corpo = linhas[1:-2]
    assert [(linha[3], linha[4], linha[5], linha[6]) for linha in corpo] == [
        ("Aluguel semana 1", recebimento.get_forma_display(), 650.0, "Sim"),
        ('=HIPERLINK("x")', "Desconto na caução", 120.0, "Não"),
    ]
    total, no_das = linhas[-2][5], linhas[-1][5]
    assert (total, no_das) == (770.0, 650.0)

    das = {linha[0]: linha[1] for linha in wb["DAS"].iter_rows(min_row=2, values_only=True)}
    assert das["Receita de locação (fatura)"] == no_das
    diversos = sum(v for k, v in das.items() if k.startswith("Pagamentos diversos"))
    assert das["Receita de locação (fatura)"] + diversos == total

    # descrição digitada que parece fórmula fica como texto na planilha
    celula = wb["Pagamentos do DAS"].cell(row=3, column=4)
    assert celula.data_type == "s"
