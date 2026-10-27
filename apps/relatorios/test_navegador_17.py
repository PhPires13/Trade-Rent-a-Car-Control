"""Etapa 17 (navegador) — relatórios: aba "Pagamentos do DAS" com sobra e caução,
data da proteção, "Já voltou" negativo, motivos da recomendação, seletor de mês
enxuto e recebíveis com vencido × a vencer e crédito."""

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.alocacoes.models import Alocacao
from apps.financeiro.models import (
    AplicacaoRecebimento,
    Caucao,
    Cobranca,
    MovimentacaoCaucao,
    MovimentoCredito,
    Recebimento,
)
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente
from apps.relatorios import services, views


@pytest.fixture
def alocacao(db):
    veiculo = Veiculo.objects.create(placa="QXQ7R28", marca_modelo="Gol")
    cliente = Cliente.objects.create(nome="Arlen", cpf_cnpj="111.222.333-44")
    return Alocacao.objects.create(
        veiculo=veiculo,
        cliente=cliente,
        data_inicio=date(2026, 8, 1),
        valor_semanal=Decimal("650.00"),
        km_entrega=10_000,
    )


def _baixar(usuario_logado, **params):
    resposta = usuario_logado.get("/relatorios/", params)
    assert resposta.status_code == 200
    return load_workbook(io.BytesIO(resposta.content))


def test_pagamentos_do_das_trazem_sobra_e_caucao_fora_do_das(usuario_logado, alocacao):
    cliente = alocacao.cliente
    aluguel = Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel semana 1",
        valor=Decimal("650.00"),
        vencimento=date(2026, 9, 8),
    )
    recebimento = Recebimento.objects.create(
        cliente=cliente, data=date(2026, 9, 8), valor=Decimal("700.00")
    )
    AplicacaoRecebimento.objects.create(
        recebimento=recebimento, cobranca=aluguel, valor=Decimal("650.00")
    )
    MovimentoCredito.objects.create(
        cliente=cliente,
        tipo=MovimentoCredito.Tipo.ENTRADA,
        valor=Decimal("50.00"),
        data=date(2026, 9, 8),
        recebimento=recebimento,
    )
    caucao = Caucao.objects.create(alocacao=alocacao)
    MovimentacaoCaucao.objects.create(
        caucao=caucao,
        tipo=MovimentacaoCaucao.Tipo.RECEBIMENTO,
        valor=Decimal("500.00"),
        data=date(2026, 9, 1),
    )

    wb = _baixar(usuario_logado, mes="2026-09", exportar="tudo")
    linhas = list(wb["Pagamentos do DAS"].iter_rows(min_row=2, values_only=True))
    fora = [linha for linha in linhas if linha[0] and linha[6] == "Não"]
    assert sorted((linha[5], linha[2]) for linha in fora) == [(50.0, None), (500.0, "QXQ7R28")]
    assert any("crédito" in linha[3] for linha in fora)
    assert any("Caução recebida" in linha[3] for linha in fora)

    # fecha com a aba "DAS"
    das = {linha[0]: linha[1] for linha in wb["DAS"].iter_rows(min_row=2, values_only=True)}
    assert das["Caução recebida (não é receita — é do cliente)"] == 500.0
    assert das["Adiantamentos recebidos (crédito do cliente)"] == 50.0


def test_mensalidade_da_protecao_tem_data_na_planilha(db):
    Veiculo.objects.create(
        placa="PRO1T00",
        marca_modelo="Onix",
        data_aquisicao=date(2026, 1, 10),
        mensalidade_protecao=Decimal("180.00"),
    )
    _, colunas, linhas = views._planilha_despesas(2026, 9)
    protecao = next(linha for linha in linhas if linha[2].startswith("Proteção"))
    assert protecao[0] == "01/09/2026"
    assert protecao[3]


def test_ja_voltou_negativo_vira_texto(db):
    assert views._ja_voltou(None) is None
    assert views._ja_voltou(Decimal("-0.09")) == "ainda não voltou"
    assert views._ja_voltou(Decimal("0.756")) == 75.6


def test_tela_e_planilha_mostram_motivo_e_nao_mostram_percentual_negativo(
    usuario_logado, monkeypatch
):
    veiculo = Veiculo.objects.create(placa="NEG1A00", marca_modelo="Gol")
    fichas, media = services.resumo_da_frota()
    ficha = next(f for f in fichas if f.veiculo.pk == veiculo.pk)
    monkeypatch.setattr(
        type(ficha), "percentual_recuperado", property(lambda self: Decimal("-0.09"))
    )
    ficha.motivos = ["5 manutenções esporádicas pesadas nos últimos 6 meses"]
    monkeypatch.setattr(views.services, "resumo_da_frota", lambda: ([ficha], media))

    html = usuario_logado.get("/relatorios/").content.decode()
    assert "ainda não voltou" in html
    assert "-9%" not in html
    assert "5 manutenções esporádicas pesadas" in html

    _, colunas, linhas = views._planilha_frota(2026, 10)
    assert linhas[0][colunas.index("Já voltou (%)")] == "ainda não voltou"
    assert "esporádicas" in linhas[0][colunas.index("Motivo")]


def test_ir_para_lista_so_meses_desde_o_primeiro_dado(db):
    hoje = date(2026, 10, 27)
    # sem nenhum dado: só o mês atual
    assert [o["param"] for o in views._opcoes_de_mes(2026, 10, hoje)] == ["2026-10"]
    Veiculo.objects.create(placa="ANT1G00", marca_modelo="Gol", data_aquisicao=date(2026, 7, 5))
    opcoes = [o["param"] for o in views._opcoes_de_mes(2026, 10, hoje)]
    assert opcoes == ["2026-10", "2026-09", "2026-08", "2026-07"]


def test_recebiveis_separam_vencido_a_vencer_e_descontam_credito(usuario_logado, alocacao):
    hoje = timezone.localdate()
    cliente = alocacao.cliente
    Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel vencido",
        valor=Decimal("320.00"),
        vencimento=hoje - timedelta(days=3),
        status=Cobranca.Status.ATRASADO,
    )
    Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel a vencer",
        valor=Decimal("650.00"),
        vencimento=hoje + timedelta(days=4),
    )
    MovimentoCredito.objects.create(
        cliente=cliente, tipo=MovimentoCredito.Tipo.ENTRADA, valor=Decimal("100.00"), data=hoje
    )

    [registro] = services.recebiveis_em_aberto()
    assert registro["vencido"] == Decimal("320.00")
    assert registro["a_vencer"] == Decimal("650.00")
    assert registro["credito"] == Decimal("100.00")
    assert registro["liquido"] == Decimal("870.00")

    html = usuario_logado.get("/relatorios/").content.decode()
    assert "deve hoje R$ 320,00" in html
    assert "R$ 650,00 a vencer" in html
    assert "R$ 100,00 de crédito" in html

    wb = _baixar(usuario_logado, exportar="recebiveis", formato="xlsx")
    ws = wb.active
    cabecalho = [c.value for c in ws[1]]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    situacoes = {linha[1]: linha[cabecalho.index("Situação")] for linha in linhas}
    assert situacoes["Aluguel vencido"] == "Vencida"
    assert situacoes["Aluguel a vencer"] == "A vencer"
    assert sum(linha[-1] for linha in linhas) == 870.0


# --- revisão: casos de borda ---------------------------------------------------


def test_ir_para_conta_aluguel_de_carro_sem_data_de_compra(alocacao):
    # carro cadastrado sem data de compra: o aluguel mais antigo vale como início
    opcoes = [o["param"] for o in views._opcoes_de_mes(2026, 10, date(2026, 10, 27))]
    assert opcoes[-1] == "2026-08"


def test_ir_para_mantem_mes_escolhido_fora_da_lista(db):
    opcoes = [o["param"] for o in views._opcoes_de_mes(2025, 3, date(2026, 10, 27))]
    assert opcoes == ["2026-10", "2025-03"]


def test_planilha_recebiveis_mostra_vence_hoje(alocacao):
    hoje = timezone.localdate()
    Cobranca.objects.create(
        cliente=alocacao.cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel de hoje",
        valor=Decimal("650.00"),
        vencimento=hoje,
    )
    [registro] = services.recebiveis_em_aberto(hoje)
    assert registro["vencido"] == Decimal("650.00")  # igual ao Quem deve
    _, colunas, linhas = views._planilha_recebiveis(hoje.year, hoje.month)
    assert linhas[0][colunas.index("Situação")] == "Vence hoje"


def test_credito_maior_que_a_divida_nao_mostra_falta_zero(usuario_logado, alocacao):
    hoje = timezone.localdate()
    Cobranca.objects.create(
        cliente=alocacao.cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao="Aluguel",
        valor=Decimal("100.00"),
        vencimento=hoje - timedelta(days=1),
    )
    MovimentoCredito.objects.create(
        cliente=alocacao.cliente,
        tipo=MovimentoCredito.Tipo.ENTRADA,
        valor=Decimal("300.00"),
        data=hoje,
    )
    [registro] = services.recebiveis_em_aberto(hoje)
    assert registro["liquido"] == Decimal("0")
    html = usuario_logado.get("/relatorios/").content.decode()
    assert "o crédito já cobre tudo" in html
    assert "falta R$ 0,00" not in html


def test_entradas_fora_do_das_sem_consulta_por_linha(alocacao, django_assert_max_num_queries):
    caucao = Caucao.objects.create(alocacao=alocacao)
    for dia in range(1, 6):
        MovimentacaoCaucao.objects.create(
            caucao=caucao,
            tipo=MovimentacaoCaucao.Tipo.REFORCO,
            valor=Decimal("10.00"),
            data=date(2026, 9, dia),
        )
        MovimentoCredito.objects.create(
            cliente=alocacao.cliente,
            tipo=MovimentoCredito.Tipo.ENTRADA,
            valor=Decimal("5.00"),
            data=date(2026, 9, dia),
        )
    with django_assert_max_num_queries(2):
        linhas = services.entradas_fora_do_das(2026, 9)
    assert len(linhas) == 10
    assert sum(linha["valor"] for linha in linhas) == Decimal("75.00")
