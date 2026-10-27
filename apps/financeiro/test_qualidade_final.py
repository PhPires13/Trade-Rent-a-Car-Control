"""Qualidade final do financeiro: cada correção da rodada "zero bugs" tem um teste aqui.

Datas relativas a hoje (HOJE): o Receber agora recusa data no futuro e pede
confirmação para pagamento de mês já fechado no DAS.
"""

import datetime as dt
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from apps.alocacoes.models import Alocacao, TrocaTemporaria
from apps.financeiro import devedores, periodos, services
from apps.financeiro.models import (
    Caucao,
    Cobranca,
    MovimentacaoCaucao,
    MovimentoCredito,
    NotaDebito,
    Recebimento,
)
from apps.frota.models import Veiculo
from apps.pessoas.models import Cliente

HOJE = timezone.localdate()
URL_RECEBER = "/financeiro/receber/"


@pytest.fixture
def cliente(db):
    return Cliente.objects.create(
        nome="Patrícia", cpf_cnpj="111.222.333-44", telefone="31987654321"
    )


@pytest.fixture
def alocacao(cliente):
    veiculo = Veiculo.objects.create(placa="RTC3D06", marca_modelo="Gol")
    return Alocacao.objects.create(
        veiculo=veiculo,
        cliente=cliente,
        data_inicio=HOJE - timedelta(days=30),
        valor_semanal=Decimal("650.00"),
        km_entrega=0,
        caucao_valor=Decimal("1000.00"),
    )


def _aluguel(cliente, alocacao=None, valor="650.00", vencimento=None, status=None):
    vencimento = vencimento or HOJE - timedelta(days=10)
    cobranca = Cobranca.objects.create(
        cliente=cliente,
        alocacao=alocacao,
        origem=Cobranca.Origem.ALUGUEL,
        descricao=f"Aluguel semanal RTC3D06 — venc. {vencimento:%d/%m/%Y}",
        valor=Decimal(valor),
        vencimento=vencimento,
        status=status or Cobranca.Status.PENDENTE,
    )
    cobranca.atualizar_status()
    return cobranca


def _pagar(cliente, cobranca, valor="650.00", data=None, **extra):
    return services.registrar_recebimento(
        cliente, data or HOJE, Decimal(valor), "pix", [(cobranca, Decimal(valor))], **extra
    )


def _post_receber(client, cliente, **campos):
    dados = {"cliente": cliente.pk, "data": HOJE.isoformat(), "forma": "pix"}
    dados.update(campos)
    return client.post(URL_RECEBER, dados)


def _mensagens(resposta):
    return [str(m) for m in resposta.context["messages"]]


# ---------- 1. pagamento numa cobrança que já foi quitada (tela velha) ----------


def test_receber_recusa_aplicacao_em_cobranca_ja_quitada(usuario_logado, cliente):
    semana = _aluguel(cliente)
    _pagar(cliente, semana)  # quitada em outra aba, há 10 minutos (não é toque duplo)
    Recebimento.history.update(history_date=timezone.now() - timedelta(minutes=10))
    resposta = _post_receber(
        usuario_logado, cliente, valor="650,00", **{f"aplicar_{semana.pk}": "650,00"}
    )
    assert resposta.status_code == 200
    assert any("já não está em aberto" in m for m in _mensagens(resposta))
    assert Recebimento.objects.count() == 1  # só o primeiro
    assert MovimentoCredito.saldo_do_cliente(cliente) == Decimal("0.00")  # nada virou crédito
    assert 'value="650,00"' in resposta.content.decode()  # o digitado continua na tela


# ---------- 2. saldo anotado velho não paga além do valor ----------


def test_saldo_velho_nao_paga_duas_vezes(cliente):
    semana = _aluguel(cliente)
    velha = Cobranca.com_quitacao_anotada(Cobranca.objects.filter(pk=semana.pk)).get()
    assert velha.saldo == Decimal("650.00")
    _pagar(cliente, semana)
    with pytest.raises(ValidationError, match="já não está em aberto"):
        services.registrar_recebimento(
            cliente, HOJE, Decimal("600"), "pix", [(velha, Decimal("600"))]
        )
    semana.refresh_from_db()
    assert semana.total_quitado == Decimal("650.00")


def test_desconto_de_caucao_com_saldo_velho(cliente, alocacao):
    caucao = services.abrir_caucao(alocacao, Decimal("1000.00"), data=alocacao.data_inicio)
    semana = _aluguel(cliente, alocacao)
    velha = Cobranca.com_quitacao_anotada(Cobranca.objects.filter(pk=semana.pk)).get()
    _pagar(cliente, semana)
    with pytest.raises(ValidationError):
        services.descontar_da_caucao(caucao, velha, Decimal("600"), HOJE)
    assert caucao.descontado == Decimal("0.00")


# ---------- 3. ND com item pela metade ----------


def _post_nd(client, cliente, descricoes, valores):
    return client.post(
        "/financeiro/nd/nova/",
        {
            "cliente": cliente.pk,
            "data_emissao": HOJE.isoformat(),
            "item_descricao": descricoes,
            "item_valor": valores,
            "token": "t-" + "-".join(descricoes),
        },
    )


def test_nd_item_com_valor_sem_descricao_e_recusado(usuario_logado, cliente):
    resposta = _post_nd(usuario_logado, cliente, ["Avaria", ""], ["350,00", "100,00"])
    assert resposta.status_code == 200
    assert any(
        "Item 2 tem valor (R$ 100,00) mas não tem descrição" in m for m in _mensagens(resposta)
    )
    assert not NotaDebito.objects.exists()
    assert resposta.context["itens_iniciais"] == [
        {"d": "Avaria", "v": "350,00"},
        {"d": "", "v": "100,00"},
    ]


def test_nd_item_sem_valor_e_recusado(usuario_logado, cliente):
    resposta = _post_nd(usuario_logado, cliente, ["Avaria", "Guincho"], ["350,00", ""])
    assert any('"Guincho" está sem valor.' in m for m in _mensagens(resposta))
    assert not NotaDebito.objects.exists()


def test_nd_linha_totalmente_vazia_so_some(usuario_logado, cliente):
    resposta = _post_nd(usuario_logado, cliente, ["Avaria", ""], ["350,00", ""])
    assert resposta.status_code == 302
    assert NotaDebito.objects.get().total == Decimal("350.00")


def test_nd_tela_soma_so_linhas_completas_e_trava_o_emitir(usuario_logado, cliente):
    html = usuario_logado.get("/financeiro/nd/nova/").content.decode()
    assert "temIncompleta" in html
    assert "filter((it) => this.completa(it))" in html


# ---------- 4. datas impossíveis no pagamento e na caução ----------


@pytest.mark.parametrize(
    "data,trecho",
    [
        (HOJE + timedelta(days=1), "não pode ser no futuro"),
        (date(26, 10, 20), "Confira o ano"),
        (date(1926, 10, 20), "Confira o ano"),
    ],
)
def test_data_do_pagamento_impossivel_e_recusada(cliente, data, trecho):
    semana = _aluguel(cliente)
    with pytest.raises(ValidationError, match=trecho):
        _pagar(cliente, semana, data=data)
    assert not Recebimento.objects.exists()


def test_pagamento_de_mes_fechado_exige_confirmacao(usuario_logado, cliente):
    semana = _aluguel(cliente)
    antiga = services.inicio_do_mes_passado(HOJE) - timedelta(days=1)
    dados = {"valor": "650,00", "data": antiga.isoformat(), f"aplicar_{semana.pk}": "650,00"}
    resposta = _post_receber(usuario_logado, cliente, **dados)
    assert resposta.status_code == 200
    assert any("Confirmo: pagamento de mês passado" in m for m in _mensagens(resposta))
    resposta = _post_receber(usuario_logado, cliente, confirmar_mes_passado="1", **dados)
    assert resposta.status_code == 302
    assert Recebimento.objects.get().data == antiga


def test_tela_receber_limita_a_data(usuario_logado, cliente):
    _aluguel(cliente)
    html = usuario_logado.get(f"{URL_RECEBER}?cliente={cliente.pk}").content.decode()
    assert f'max="{HOJE.isoformat()}"' in html
    assert 'name="confirmar_mes_passado"' in html


def test_caucao_recusa_data_futura_e_antes_do_contrato(cliente, alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    with pytest.raises(ValidationError, match="futuro"):
        services.registrar_movimentacao_caucao(
            caucao, "reforco", Decimal("100"), HOJE + timedelta(days=200), forma="pix"
        )
    services.registrar_movimentacao_caucao(
        caucao, "recebimento", Decimal("1000"), alocacao.data_inicio, forma="pix"
    )
    with pytest.raises(ValidationError, match="antes do início do contrato"):
        services.registrar_movimentacao_caucao(
            caucao,
            "devolucao",
            Decimal("100"),
            alocacao.data_inicio - timedelta(days=900),
            devolver_mesmo_assim=True,
        )


# ---------- 5. desconto de caução numa cobrança cancelada ----------


def test_desconto_de_caucao_em_cobranca_cancelada_e_recusado(usuario_logado, cliente, alocacao):
    caucao = services.abrir_caucao(alocacao, Decimal("1000.00"), data=alocacao.data_inicio)
    semana = _aluguel(cliente, alocacao, status=Cobranca.Status.CANCELADA)
    semana.status = Cobranca.Status.CANCELADA
    semana.save(update_fields=["status"])
    with pytest.raises(ValidationError, match="não está mais em aberto"):
        services.descontar_da_caucao(caucao, semana, Decimal("300"), HOJE)
    resposta = usuario_logado.post(
        f"/financeiro/caucoes/{caucao.pk}/",
        {
            "tipo": "desconto",
            "valor": "300,00",
            "data": HOJE.isoformat(),
            "cobranca": semana.pk,
            "token": "x1",
        },
    )
    assert resposta.status_code == 200
    assert caucao.saldo == Decimal("1000.00")


# ---------- 6. dia de vencimento mudado no meio da semana (D04) ----------


@pytest.fixture
def alocacao_julho(cliente):
    veiculo = Veiculo.objects.create(placa="QXQ6C10", marca_modelo="Gol")
    return Alocacao.objects.create(
        veiculo=veiculo,
        cliente=cliente,
        data_inicio=date(2026, 7, 1),  # quarta
        valor_semanal=Decimal("650.00"),
        km_entrega=0,
    )


def test_mudanca_do_dia_nao_cria_cobranca_ja_vencida(cliente, alocacao_julho):
    alocacao = alocacao_julho
    for cobranca in services.gerar_cobrancas_semanais(hoje=date(2026, 7, 16)):  # 01, 08 e 15/07
        services.registrar_recebimento(
            cliente,
            date(2026, 7, 16),
            cobranca.valor,
            "pix",
            [(cobranca, cobranca.valor)],
            hoje=date(2026, 7, 16),
        )
    alocacao.dia_vencimento = 4  # sexta, mudado na segunda 20/07
    alocacao.save()
    alocacao.history.exclude(dia_vencimento=4).update(
        history_date=dt.datetime(2026, 7, 1, 12, tzinfo=dt.UTC)
    )
    alocacao.history.filter(dia_vencimento=4).update(
        history_date=dt.datetime(2026, 7, 20, 12, tzinfo=dt.UTC)
    )
    assert services.gerar_cobrancas_semanais(hoje=date(2026, 7, 21)) == []
    services.marcar_atrasos(hoje=date(2026, 7, 21))
    cliente.refresh_from_db()
    assert cliente.status == Cliente.Status.ATIVO
    assert services.proximo_vencimento(alocacao, date(2026, 7, 21)) == (
        date(2026, 7, 24),
        date(2026, 7, 22),
    )
    (transicao,) = services.gerar_cobrancas_semanais(hoje=date(2026, 7, 24))
    assert transicao.vencimento == date(2026, 7, 24)
    assert services.periodo_da_cobranca(transicao) == (date(2026, 7, 22), date(2026, 7, 30))
    assert transicao.valor == Decimal("835.71")  # 650 / 7 × 9
    assert "9 dias" in transicao.descricao


# ---------- 7. valor semanal corrigido refaz a semana em aberto ----------


def test_refazer_alugueis_em_aberto_corrige_o_valor_digitado_errado(cliente, alocacao_julho):
    alocacao = alocacao_julho
    alocacao.valor_semanal = Decimal("6500.00")
    alocacao.save()
    (semana,) = services.gerar_cobrancas_semanais(hoje=date(2026, 7, 1))
    assert semana.valor == Decimal("6500.00")
    alocacao.valor_semanal = Decimal("650.00")
    alocacao.save()
    assert services.refazer_alugueis_em_aberto(alocacao, hoje=date(2026, 7, 1)) == [semana]
    semana.refresh_from_db()
    assert semana.valor == Decimal("650.00")


def test_refazer_nao_baixa_abaixo_do_que_ja_foi_pago(cliente, alocacao_julho):
    alocacao = alocacao_julho
    (semana,) = services.gerar_cobrancas_semanais(hoje=date(2026, 7, 1))
    services.registrar_recebimento(
        cliente,
        date(2026, 7, 1),
        Decimal("600"),
        "pix",
        [(semana, Decimal("600"))],
        hoje=date(2026, 7, 1),
    )
    alocacao.valor_semanal = Decimal("500.00")
    alocacao.save()
    assert services.refazer_alugueis_em_aberto(alocacao, hoje=date(2026, 7, 1)) == []
    semana.refresh_from_db()
    assert semana.valor == Decimal("650.00")


# ---------- 8. encargo só em atrasada e com teto ----------


def test_encargo_recusado_em_cobranca_paga(cliente):
    semana = _aluguel(cliente)
    _pagar(cliente, semana)
    with pytest.raises(ValidationError, match="não está atrasada"):
        services.aplicar_encargo(semana, Decimal("32.50"))


def test_encargo_recusado_em_cobranca_que_ainda_vai_vencer(cliente):
    semana = _aluguel(cliente, vencimento=HOJE + timedelta(days=3))
    with pytest.raises(ValidationError, match="não está atrasada"):
        services.aplicar_encargo(semana, Decimal("32.50"))


def test_encargo_acima_de_10_por_cento_e_recusado(cliente):
    semana = _aluguel(cliente)
    with pytest.raises(ValidationError, match="R\\$ 65,00"):
        services.aplicar_encargo(semana, Decimal("650.00"))
    assert services.aplicar_encargo(semana, Decimal("65.00")).valor == Decimal("65.00")


def test_encargo_aceita_porcentagem(usuario_logado, cliente):
    semana = _aluguel(cliente)
    usuario_logado.post(f"/financeiro/cobranca/{semana.pk}/encargo/", {"valor": "5%"})
    assert semana.encargos.get().valor == Decimal("32.50")


# ---------- 9. encargo cobrado junto no Receber ----------


def test_receber_inclui_o_encargo_no_mesmo_lancamento(usuario_logado, cliente):
    semana = _aluguel(cliente)  # 10 dias de atraso → 10% sugerido
    html = usuario_logado.get(f"{URL_RECEBER}?cliente={cliente.pk}").content.decode()
    assert f'name="incluir_encargo_{semana.pk}"' in html
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="715,00",
        **{
            f"aplicar_{semana.pk}": "650,00",
            f"incluir_encargo_{semana.pk}": "1",
            f"encargo_{semana.pk}": "65,00",
            f"aplicar_encargo_{semana.pk}": "65,00",
        },
    )
    assert resposta.status_code == 302
    encargo = semana.encargos.get()
    assert encargo.valor == Decimal("65.00")
    encargo.refresh_from_db()
    assert encargo.status == Cobranca.Status.PAGO
    assert MovimentoCredito.saldo_do_cliente(cliente) == Decimal("0.00")


def test_desfazer_o_lancamento_cancela_o_encargo_cobrado_junto(usuario_logado, cliente):
    semana = _aluguel(cliente)
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="715,00",
        **{
            f"aplicar_{semana.pk}": "650,00",
            f"incluir_encargo_{semana.pk}": "1",
            f"encargo_{semana.pk}": "65,00",
            f"aplicar_encargo_{semana.pk}": "65,00",
        },
    )
    recebimento = Recebimento.objects.get()
    assert resposta.status_code == 302
    usuario_logado.post(f"/financeiro/recebimento/{recebimento.pk}/desfazer/")
    assert semana.encargos.get().status == Cobranca.Status.CANCELADA


def test_encargo_some_se_o_lancamento_falha(usuario_logado, cliente):
    semana = _aluguel(cliente)
    _post_receber(
        usuario_logado,
        cliente,
        valor="100,00",  # menos do que foi distribuído → erro
        **{
            f"aplicar_{semana.pk}": "650,00",
            f"incluir_encargo_{semana.pk}": "1",
            f"encargo_{semana.pk}": "65,00",
        },
    )
    assert not semana.encargos.exists()


# ---------- 10. caução: desfazer, combinado, acima do combinado, observações ----------


def test_caucao_mostra_combinado_e_preenche_o_que_falta(usuario_logado, alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    resposta = usuario_logado.get(f"/financeiro/caucoes/{caucao.pk}/")
    html = resposta.content.decode()
    assert "Combinada" in html and "falta R$ 1.000,00" in html
    assert 'value="1.000,00"' in html
    assert "— Trade Rent a Car</title>" in html


def test_caucao_acima_do_combinado_exige_confirmacao(usuario_logado, alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    dados = {"tipo": "recebimento", "valor": "10.000,00", "data": HOJE.isoformat(), "forma": "pix"}
    resposta = usuario_logado.post(f"/financeiro/caucoes/{caucao.pk}/", {**dados, "token": "a"})
    assert resposta.status_code == 200
    assert any("acima do" in m for m in _mensagens(resposta))
    assert caucao.recebido == Decimal("0.00")
    resposta = usuario_logado.post(
        f"/financeiro/caucoes/{caucao.pk}/", {**dados, "token": "b", "confirmar_acima": "1"}
    )
    assert resposta.status_code == 302
    assert caucao.recebido == Decimal("10000.00")


def test_caucao_desfazer_lancamento_do_dia(usuario_logado, alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    movimentacao = services.registrar_movimentacao_caucao(
        caucao, "recebimento", Decimal("1000"), HOJE, forma="pix", observacoes="depósito\x00 Pix"
    )
    assert movimentacao.observacoes == "depósito Pix"  # NUL some (Postgres recusaria)
    html = usuario_logado.get(f"/financeiro/caucoes/{caucao.pk}/").content.decode()
    assert "↩ desfazer" in html and "depósito Pix" in html
    usuario_logado.post(f"/financeiro/caucoes/movimentacao/{movimentacao.pk}/desfazer/")
    assert not MovimentacaoCaucao.objects.exists()


def test_caucao_desfazer_de_outro_dia_e_recusado(alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    movimentacao = services.registrar_movimentacao_caucao(
        caucao, "recebimento", Decimal("1000"), HOJE, forma="pix"
    )
    with pytest.raises(ValidationError, match="mesmo dia"):
        services.desfazer_movimentacao_caucao(movimentacao, hoje=HOJE + timedelta(days=1))


def test_caucao_forma_fora_da_lista_e_recusada(alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    with pytest.raises(ValidationError, match="Escolha a forma"):
        services.registrar_movimentacao_caucao(
            caucao, "recebimento", Decimal("100"), HOJE, forma="cartao-de-credito-parcelado"
        )


# ---------- 11. devolver a caução com dívida / carro com o cliente ----------


def test_devolver_caucao_com_divida_exige_confirmacao(usuario_logado, cliente, alocacao):
    caucao = Caucao.objects.get(alocacao=alocacao)
    services.registrar_movimentacao_caucao(caucao, "recebimento", Decimal("1000"), HOJE, "pix")
    _aluguel(cliente, alocacao, valor="428.57")
    html = usuario_logado.get(f"/financeiro/caucoes/{caucao.pk}/").content.decode()
    assert "Patrícia ainda deve R$ 428,57" in html and "RTC3D06" in html
    assert "Usar a caução para pagar" in html
    with pytest.raises(ValidationError, match="Devolver mesmo assim"):
        services.registrar_movimentacao_caucao(caucao, "devolucao", Decimal("1000"), HOJE)
    services.registrar_movimentacao_caucao(
        caucao, "devolucao", Decimal("1000"), HOJE, devolver_mesmo_assim=True
    )
    assert caucao.devolvido == Decimal("1000.00")


# ---------- 12. crédito + dinheiro novo num lançamento só ----------


def _credito(cliente, valor="50.00"):
    MovimentoCredito.objects.create(
        cliente=cliente, tipo=MovimentoCredito.Tipo.ENTRADA, valor=Decimal(valor), data=HOJE
    )


def test_credito_e_pix_juntos_quitam_a_semana(usuario_logado, cliente):
    semana = _aluguel(cliente)
    _credito(cliente)
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="600,00",
        usar_credito_junto="1",
        **{f"aplicar_{semana.pk}": "650,00"},
    )
    assert resposta.status_code == 302
    semana.refresh_from_db()
    assert semana.status == Cobranca.Status.PAGO
    assert MovimentoCredito.saldo_do_cliente(cliente) == Decimal("0.00")
    formas = sorted(Recebimento.objects.values_list("forma", "valor"))
    assert formas == [("credito", Decimal("50.00")), ("pix", Decimal("600.00"))]
    fiscal = services.resumo_fiscal(HOJE.year, HOJE.month)
    assert fiscal["locacao"] == Decimal("650.00")
    assert fiscal["pago_com_credito"] == Decimal("50.00")
    # o resumo mostra os dois e o desfazer leva os dois juntos
    html = usuario_logado.get(resposta.url).content.decode()
    assert "Crédito usado junto" in html
    principal = Recebimento.objects.get(forma="pix")
    junto = Recebimento.objects.get(forma="credito")
    usuario_logado.post(f"/financeiro/recebimento/{principal.pk}/desfazer/", {"junto": junto.pk})
    assert not Recebimento.objects.exists()
    assert MovimentoCredito.saldo_do_cliente(cliente) == Decimal("50.00")


def test_so_credito_mostra_mensagem_de_credito(usuario_logado, cliente):
    semana = _aluguel(cliente)
    _credito(cliente, "650.00")
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="650,00",
        forma="credito",
        **{f"aplicar_{semana.pk}": "650,00"},
    )
    mensagens = [str(m) for m in usuario_logado.get(resposta.url).context["messages"]]
    assert any(m.startswith("Crédito de R$ 650,00 usado para Patrícia") for m in mensagens)


# ---------- 13 e 14. Quem deve: crédito parado e texto do "vence hoje" ----------


def test_quem_deve_mostra_credito_parado_com_atalho(usuario_logado, cliente):
    _aluguel(cliente, valor="145.00")
    _credito(cliente, "200.00")
    (devedor,) = devedores.devedores()
    assert devedor.credito == Decimal("200.00")
    html = usuario_logado.get("/financeiro/quem-deve/").content.decode()
    assert "de crédito" in html
    assert f"/financeiro/receber/?cliente={cliente.pk}&usar_credito=1" in html
    receber = usuario_logado.get(f"{URL_RECEBER}?cliente={cliente.pk}&usar_credito=1")
    assert "usarCredito();" in receber.content.decode()


def test_vence_hoje_nao_parece_somar_ao_atrasado(usuario_logado, cliente):
    _aluguel(cliente, valor="8250.00", vencimento=HOJE - timedelta(days=91))
    _aluguel(cliente, valor="650.00", vencimento=HOJE)
    html = usuario_logado.get("/financeiro/quem-deve/").content.decode()
    assert "R$ 8.900,00" in html
    assert "+ R$ 650,00 vence hoje" not in html
    assert "atrasado há 91 dias" in html and "e <span" in html


def test_vencido_ate_hoje_e_funcao_unica(cliente):
    _aluguel(cliente, valor="100.00", vencimento=HOJE - timedelta(days=3))
    _aluguel(cliente, valor="65.00", vencimento=HOJE + timedelta(days=3))
    judicial = _aluguel(cliente, valor="40.00", vencimento=HOJE - timedelta(days=30))
    judicial.status = Cobranca.Status.JUDICIAL
    judicial.save(update_fields=["status"])
    assert devedores.vencido_ate_hoje([cliente.pk]) == {cliente.pk: Decimal("140.00")}
    assert devedores.vencido_ate_hoje(incluir_judicial=False) == {cliente.pk: Decimal("100.00")}


# ---------- 15. Receber separa vencido de a vencer; filtro "Em aberto" mostra judicial ----------


def test_receber_separa_vencido_de_a_vencer(usuario_logado, cliente):
    _aluguel(cliente, valor="100.00", vencimento=HOJE - timedelta(days=3))
    _aluguel(cliente, valor="65.00", vencimento=HOJE + timedelta(days=3))
    resposta = usuario_logado.get(f"{URL_RECEBER}?cliente={cliente.pk}")
    assert resposta.context["vencido"] == Decimal("100.00")
    assert resposta.context["a_vencer"] == Decimal("65.00")
    html = resposta.content.decode()
    assert "Vencido:" in html and "A vencer:" in html


def test_filtro_em_aberto_mostra_judicial(usuario_logado, cliente):
    judicial = _aluguel(cliente)
    judicial.status = Cobranca.Status.JUDICIAL
    judicial.save(update_fields=["status"])
    html = usuario_logado.get("/financeiro/").content.decode()
    assert f'id="cobranca-{judicial.pk}"' in html
    assert "Na Justiça" in html


# ---------- 16. Extrato com totais no topo ----------


def test_extrato_mostra_totais_e_falta(usuario_logado, cliente):
    semana = _aluguel(cliente)
    services.registrar_recebimento(cliente, HOJE, Decimal("380"), "pix", [(semana, Decimal("330"))])
    html = usuario_logado.get(f"/financeiro/cliente/{cliente.pk}/").content.decode()
    assert "Deve hoje" in html and "A vencer" in html and "Caução guardada" in html
    assert "falta R$ 320,00" in html
    assert "sobra" in html and "→ crédito" in html
    assert "Extrato — Patrícia — Trade Rent a Car" in html


# ---------- 17 e 18. DAS: planilha com 2 abas, mês padrão e setas ----------


def test_planilha_do_das_tem_aba_de_pagamentos_que_bate(usuario_logado, cliente):
    semana = _aluguel(cliente)
    _pagar(cliente, semana)
    mes = f"{HOJE.year}-{HOJE.month:02d}"
    resposta = usuario_logado.get(f"/financeiro/das/?mes={mes}&exportar=xlsx")
    livro = load_workbook(BytesIO(resposta.content))
    assert livro.sheetnames == ["Base do DAS", "Pagamentos do mês"]
    pagamentos = list(livro["Pagamentos do mês"].iter_rows(values_only=True))
    assert pagamentos[1][1] == "Patrícia" and pagamentos[1][6] == "Sim"
    total = next(linha for linha in pagamentos if linha[3] and linha[3].startswith("Total"))
    base = list(livro["Base do DAS"].iter_rows(values_only=True))[1]
    assert total[5] == base[1] == 650.0


@pytest.mark.parametrize(
    "hoje,esperado",
    [
        (date(2026, 11, 5), (2026, 10)),
        (date(2026, 11, 20), (2026, 10)),
        (date(2026, 11, 21), (2026, 11)),
        (date(2027, 1, 3), (2026, 12)),
    ],
)
def test_das_abre_no_mes_do_prazo(hoje, esperado):
    assert periodos.mes_padrao_do_das(hoje) == esperado


def test_tela_do_das_usa_setas_e_nao_input_month(usuario_logado, db):
    html = usuario_logado.get("/financeiro/das/?mes=2026-10").content.decode()
    assert 'type="month"' not in html
    assert "DAS de Outubro de 2026" in html
    assert "◀ mês anterior" in html


# ---------- 19. Cobranças e cauções no celular ----------


def test_cobrancas_recolhe_acoes_raras_em_mais(usuario_logado, cliente):
    _aluguel(cliente)
    html = usuario_logado.get("/financeiro/").content.decode()
    assert "<details" in html and "Mais ▾" in html
    assert "md:table" in html  # cartão no celular, tabela do md para cima


def test_acao_volta_na_linha_da_cobranca(usuario_logado, cliente):
    semana = _aluguel(cliente)
    resposta = usuario_logado.post(
        f"/financeiro/cobranca/{semana.pk}/judicial/", {"voltar": "status=abertas"}
    )
    assert resposta["Location"].endswith(f"#cobranca-{semana.pk}")


# ---------- 20. Admin financeiro sem "Adicionar" ----------


@pytest.mark.parametrize(
    "modelo", ["cobranca", "recebimento", "caucao", "notadebito", "movimentocredito"]
)
def test_admin_financeiro_nao_inclui(admin_client, modelo):
    assert admin_client.get(f"/admin/financeiro/{modelo}/add/").status_code == 403
    assert admin_client.get(f"/admin/financeiro/{modelo}/").status_code == 200


# ---------- 21. Seletor de cliente com placa e saldo ----------


def test_seletor_de_cliente_distingue_homonimos(usuario_logado, cliente, alocacao):
    _aluguel(cliente, alocacao, valor="428.57")
    Cliente.objects.create(nome="Patrícia Souza", cpf_cnpj="999.888.777-66")
    html = usuario_logado.get(URL_RECEBER).content.decode()
    assert "Patrícia — RTC3D06 — deve hoje R$ 428,57" in html
    assert 'label="Com valor em aberto"' in html and 'label="Demais clientes"' in html
    assert "Patrícia — RTC3D06" in usuario_logado.get("/financeiro/nd/nova/").content.decode()


# ---------- 22. Receber: sobra e reforço de caução ----------


def test_reforco_de_caucao_so_para_quem_tem_caucao(usuario_logado, cliente, alocacao):
    outro = Cliente.objects.create(nome="Carlos", cpf_cnpj="999.888.777-66")
    sem = usuario_logado.get(f"{URL_RECEBER}?cliente={outro.pk}").content.decode()
    assert 'value="caucao"' not in sem
    com = usuario_logado.get(f"{URL_RECEBER}?cliente={cliente.pk}").content.decode()
    assert 'value="caucao"' in com
    assert "disponivel > aplicado + 0.005" in com  # "O que sobrar vira" só quando sobra


# ---------- 23. Descrição da ND e confirmação de cancelar ----------


def test_descricao_da_cobranca_da_nd_diz_o_que_e(cliente):
    nd = services.emitir_nota_debito(
        cliente, HOJE, [("Combustível faltando", "80.00"), ("Lavagem", "40.00")]
    )
    assert nd.cobranca.descricao == f"ND {nd.numero:03d} — Combustível faltando +1 item"


def test_cancelar_nd_sem_multa_nao_fala_de_multa(usuario_logado, cliente):
    nd = services.emitir_nota_debito(cliente, HOJE, [("Avaria", "350.00")])
    html = usuario_logado.get("/financeiro/nd/").content.decode()
    assert f"Cancelar a ND {nd.numero:03d} (R$ 350,00)? Ela deixa de ser cobrada" in html
    assert "multas dela voltam" not in html


# ---------- 24. NDs com "✓ Recebi" e filtro ----------


def test_nds_filtram_por_situacao_e_tem_recebi(usuario_logado, cliente):
    aberta = services.emitir_nota_debito(cliente, HOJE, [("Avaria", "350.00")])
    paga = services.emitir_nota_debito(cliente, HOJE, [("Guincho", "100.00")])
    _pagar(cliente, paga.cobranca, "100.00")
    resposta = usuario_logado.get("/financeiro/nd/")
    assert [nd.pk for nd in resposta.context["notas"]] == [aberta.pk]
    assert f"cobranca={aberta.cobranca.pk}" in resposta.content.decode()
    pagas = usuario_logado.get("/financeiro/nd/?situacao=pagas").context["notas"]
    assert [nd.pk for nd in pagas] == [paga.pk]


# ---------- 25. Troca temporária mantém o preço que a semana já tinha ----------


def test_troca_mantem_o_preco_dos_dias_ja_cobrados(cliente, alocacao_julho):
    alocacao = alocacao_julho
    alocacao.valor_semanal = Decimal("700.00")
    alocacao.save()
    (semana,) = services.gerar_cobrancas_semanais(hoje=date(2026, 7, 1))  # 01–07/07 a 700
    alocacao.valor_semanal = Decimal("630.00")
    alocacao.save()
    substituto = Veiculo.objects.create(placa="RNB9J66", marca_modelo="Voyage")
    TrocaTemporaria.objects.create(
        alocacao=alocacao,
        veiculo_substituto=substituto,
        data_retirada=date(2026, 7, 5),  # 5º dia
        km_retirada=0,
        valor_semanal_ajustado=Decimal("700.00"),
    )
    semana.refresh_from_db()
    assert semana.valor == Decimal("700.00")  # 4 dias a 700 + 3 dias a 700, não 660


# ---------- 26. Observações e forma lidas à mão ----------


def test_observacao_com_nul_nao_quebra_o_receber(usuario_logado, cliente):
    semana = _aluguel(cliente)
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="650,00",
        observacoes="pago\x00 na loja",
        **{f"aplicar_{semana.pk}": "650,00"},
    )
    assert resposta.status_code == 302
    assert Recebimento.objects.get().observacoes == "pago na loja"


# ---------- 27. Rótulos e data local ----------


def test_rotulos_de_repasse_em_linguagem_simples():
    assert Cobranca.Origem.REPASSE_MANUTENCAO.label == "Conserto cobrado do cliente"
    assert Cobranca.Origem.REPASSE_SINISTRO.label == "Franquia de sinistro cobrada do cliente"


def test_servicos_usam_a_data_local(cliente):
    semana = _aluguel(cliente, vencimento=date(2026, 7, 1))
    semana.status = Cobranca.Status.PENDENTE
    semana.save(update_fields=["status"])
    with mock.patch("django.utils.timezone.localdate", return_value=date(2026, 6, 30)):
        semana.atualizar_status()
    assert semana.status == Cobranca.Status.PENDENTE  # no "hoje" local ainda não venceu


# ---------- revisão: desfazer crédito + dinheiro com encargo pago pelo crédito ----------


def test_desfazer_credito_e_pix_cancela_encargo_pago_pelo_credito(usuario_logado, cliente):
    """Crédito cobre o aluguel e parte do encargo: desfazer não deixa encargo órfão."""
    semana = _aluguel(cliente)  # 10 dias de atraso → 10% (65,00)
    _credito(cliente, "700.00")
    resposta = _post_receber(
        usuario_logado,
        cliente,
        valor="15,00",
        usar_credito_junto="1",
        **{
            f"aplicar_{semana.pk}": "650,00",
            f"incluir_encargo_{semana.pk}": "1",
            f"encargo_{semana.pk}": "65,00",
            f"aplicar_encargo_{semana.pk}": "65,00",
        },
    )
    assert resposta.status_code == 302
    encargo = semana.encargos.get()
    encargo.refresh_from_db()
    assert encargo.status == Cobranca.Status.PAGO
    principal = Recebimento.objects.get(forma="pix")
    junto = Recebimento.objects.get(forma="credito")
    assert junto.aplicacoes.filter(cobranca=encargo).exists()  # o crédito pagou parte do encargo
    usuario_logado.post(f"/financeiro/recebimento/{principal.pk}/desfazer/", {"junto": junto.pk})
    assert not Recebimento.objects.exists()
    encargo.refresh_from_db()
    assert encargo.status == Cobranca.Status.CANCELADA
    assert MovimentoCredito.saldo_do_cliente(cliente) == Decimal("700.00")
    semana.refresh_from_db()
    assert semana.status == Cobranca.Status.ATRASADO


def test_planilha_csv_do_das_comeca_com_bom(usuario_logado, db):
    resposta = usuario_logado.get(f"/financeiro/das/?mes={HOJE:%Y-%m}&exportar=csv")
    assert resposta.content.startswith("﻿".encode())


@pytest.mark.parametrize(
    "url",
    [
        "/financeiro/das/?mes=0001-01",
        "/financeiro/das/?mes=9999-12",
        "/financeiro/das/?mes=1999-12&exportar=xlsx",
        "/financeiro/nd/?situacao=xyz",
        "/financeiro/receber/?cliente=abc&usar_credito=1",
        "/financeiro/receber/?cliente=999999&cobranca=999999",
        "/financeiro/receber/?recibo=abc&recibo_credito=xyz",
        "/financeiro/?status=abertas&pagina=999",
        "/financeiro/?status=xyz&origem=abc",
    ],
)
def test_parametros_estranhos_nao_dao_erro_500(usuario_logado, cliente, url):
    assert usuario_logado.get(url).status_code < 500  # cliente inexistente: 404


def test_desfazer_movimentacao_inexistente_da_404(usuario_logado, db):
    resposta = usuario_logado.post("/financeiro/caucoes/movimentacao/999999/desfazer/")
    assert resposta.status_code == 404


def test_receber_com_aplicar_de_outro_cliente_e_recusado(usuario_logado, cliente):
    outro = Cliente.objects.create(nome="Outro", cpf_cnpj="999.888.777-66")
    alheia = _aluguel(outro)
    resposta = _post_receber(
        usuario_logado, cliente, valor="650,00", **{f"aplicar_{alheia.pk}": "650,00"}
    )
    assert resposta.status_code == 200
    assert not Recebimento.objects.exists()
    alheia.refresh_from_db()
    assert alheia.status == Cobranca.Status.ATRASADO
