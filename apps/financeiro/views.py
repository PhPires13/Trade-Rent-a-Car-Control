import csv
from contextlib import contextmanager
from datetime import timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from uuid import uuid4

from django.contrib import messages
from django.contrib.sessions.backends.db import SessionStore as SessaoNoBanco
from django.contrib.sessions.models import Session
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Exists, OuterRef, Prefetch, Subquery, Sum
from django.http import Http404, HttpResponse, QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme, urlencode

from apps.campos import CampoDinheiro, exibir_dinheiro, ler_data
from apps.formatos import brl
from apps.pessoas.models import Cliente

from . import recibo, services, whatsapp
from .devedores import creditos_por_cliente
from .models import (
    ZERO,
    Caucao,
    Cobranca,
    ItemNotaDebito,
    MovimentacaoCaucao,
    MovimentoCredito,
    NotaDebito,
    Recebimento,
)
from .periodos import ANO_MINIMO, ano_mes, mes_padrao_do_das, navegacao_de_meses

COBRANCAS_POR_PAGINA = 50

#: Origens que o dono pode decidir não cobrar: excedente de km, encargo por
#: atraso (decisão nº 13, "ajustável caso a caso") e cobranças avulsas.
#: A ND entra pela decisão D05: multa anulada no recurso → cancela-se a ND
#: inteira (sem pagamento nem judicial) e emite-se outra sem ela.
ORIGENS_CANCELAVEIS = [
    Cobranca.Origem.EXCEDENTE_KM,
    Cobranca.Origem.ENCARGO,
    Cobranca.Origem.OUTRO,
    Cobranca.Origem.NOTA_DEBITO,
]

CENTAVO = Decimal("0.01")
MAX_DESCRICAO_ITEM_ND = ItemNotaDebito._meta.get_field("descricao").max_length


def _dinheiro(texto, invalido):
    """Valor digitado ("1.250,00") validado como os campos de dinheiro do financeiro.

    Antes só o dinheiro_br era usado: "1e10" ou "999999999" passavam, estouravam
    o numeric(10,2) no Postgres (erro 500) e, no SQLite, derrubavam o extrato (C07).
    Arredonda para centavos como o banco faria. `invalido`: mensagem do campo.
    """
    campo = Cobranca._meta.get_field("valor").formfield(form_class=CampoDinheiro)
    try:
        valor = campo.to_python((texto or "").strip())
        if valor is None:
            raise ValidationError(invalido)
        centavos = valor.quantize(CENTAVO, rounding=ROUND_HALF_UP)
    except ValidationError as erro:
        if getattr(erro, "code", None) == "valor_ambiguo":
            raise  # "12,000": a mensagem explica o que corrigir
        raise ValidationError(invalido) from None
    except InvalidOperation:
        raise ValidationError(invalido) from None
    if centavos != valor:
        # "1.000,0010000" (digitou por cima do valor que já vinha no campo) virava
        # R$ 1.000,00 em silêncio; "650,0000" (= 650,00) continua aceito
        raise ValidationError(
            f"Valor com mais de 2 casas depois da vírgula ({texto.strip()}): "
            "use só 2 casas (ex.: 1.250,00)."
        )
    valor = centavos
    if valor < ZERO:
        # negativo na distribuição do Receber inflava a sobra (crédito que não existe)
        raise ValidationError(
            f"O valor não pode ser negativo ({texto.strip()}): use um valor maior que zero."
        )
    try:
        campo.run_validators(valor)
    except ValidationError:
        raise ValidationError(
            f"Valor alto demais ({texto.strip()}): o máximo é R$ 99.999.999,99. "
            "Confira se não sobrou algum número."
        ) from None
    return valor


def _tokens_usados(request, chave):
    """Tokens já gravados, lidos direto do banco com a linha da sessão travada.

    A sessão em memória foi lida no começo do pedido: num clique duplo com
    internet lenta, o 2º envio chega antes do 1º terminar e não veria o token.
    Travando a linha, o 2º espera o 1º gravar e então lê o token já usado.
    """
    usados = request.session.get(chave, [])
    chave_sessao = getattr(request.session, "session_key", None)
    if chave_sessao and isinstance(request.session, SessaoNoBanco):
        linha = Session.objects.select_for_update().filter(session_key=chave_sessao).first()
        if linha:
            usados = linha.get_decoded().get(chave, [])
    return usados


@contextmanager
def _envio_unico(request, chave, token):
    """Roda o lançamento uma vez só por formulário (clique duplo — C06, como na ND).

    Rende True quando o token já foi usado (nada deve ser gravado). Se o bloco
    terminar sem erro, o token fica gravado na sessão na mesma transação.
    """
    with transaction.atomic():
        usados = _tokens_usados(request, chave) if token else []
        if token and token in usados:
            yield True
            return
        yield False
        if token:
            request.session[chave] = (usados + [token])[-20:]
            if isinstance(request.session, SessaoNoBanco) and request.session.session_key:
                request.session.save()  # grava antes de soltar a trava


#: o que o painel de cobranças aceita de volta depois de uma ação (C29)
FILTROS_DO_PAINEL = ("status", "origem", "pagina")


def _voltar_ao_painel(request, cobranca=None):
    """Volta para onde a pessoa estava: filtros + página do painel, ou a lista de NDs.

    O campo "voltar" é remontado só com os parâmetros conhecidos do próprio
    painel — nunca vira outro endereço. Com a cobrança, volta na linha dela
    (âncora #cobranca-<pk>), e não no topo da lista.
    """
    voltar = request.POST.get("voltar", "")
    if voltar == "nds":
        return redirect("financeiro:nds")
    recebidos = QueryDict(voltar.lstrip("?"))
    filtros = {chave: recebidos[chave][:30] for chave in FILTROS_DO_PAINEL if recebidos.get(chave)}
    destino = reverse("financeiro:cobrancas")
    if filtros:
        destino += f"?{urlencode(filtros)}"
    # com erro, sem a âncora: a página abre no topo, onde a mensagem aparece
    # (antes rolava até a linha e o aviso de erro ficava fora da tela)
    if cobranca is not None and not getattr(request, "erro_no_painel", False):
        destino += f"#cobranca-{cobranca.pk}"
    return redirect(destino)


def _erro_no_painel(request, texto):
    """Mensagem de erro de uma ação do painel — a volta não pula para a linha."""
    request.erro_no_painel = True
    messages.error(request, texto)


def _anotar_encargos(queryset):
    """Soma e quantidade de encargos não cancelados de cada cobrança (sugestão sem N+1)."""
    encargos = (
        Cobranca.objects.filter(cobranca_origem=OuterRef("pk"))
        .exclude(status=Cobranca.Status.CANCELADA)
        .values("cobranca_origem")
    )
    return queryset.annotate(
        encargo_aplicado=Subquery(encargos.annotate(t=Sum("valor")).values("t")),
        encargos_ativos=Subquery(encargos.annotate(n=Count("pk")).values("n")),
    )


def _cobrancas_devidas(cliente):
    """O que o cliente ainda deve (judicial inclusa), com o saldo anotado em lote."""
    return list(
        _anotar_encargos(
            Cobranca.com_quitacao_anotada(
                cliente.cobrancas.filter(status__in=Cobranca.STATUS_DEVIDOS).order_by(
                    "vencimento", "pk"
                )
            )
        )
    )


def _abertas_por_cliente(cliente_ids):
    """{cliente_id: [cobranças no fluxo normal com saldo]} numa consulta só.

    Alimenta a mensagem única de WhatsApp por cliente — judicial fica de fora.
    """
    abertas = {}
    qs = Cobranca.com_quitacao_anotada(
        Cobranca.objects.filter(
            cliente_id__in=cliente_ids, status__in=Cobranca.STATUS_EM_ABERTO
        ).select_related("cliente", "alocacao__veiculo")
    ).order_by("vencimento", "pk")
    for cobranca in qs:
        if cobranca.saldo > ZERO:
            abertas.setdefault(cobranca.cliente_id, []).append(cobranca)
    return abertas


def cobrancas(request):
    """Painel de cobranças com filtros, paginação e sugestão de encargo (docs.md §4.3)."""
    encargos = (
        Cobranca.objects.filter(cobranca_origem=OuterRef("pk"))
        .exclude(status=Cobranca.Status.CANCELADA)
        .values("cobranca_origem")
    )
    qs = Cobranca.com_quitacao_anotada(
        Cobranca.objects.select_related("cliente", "alocacao__veiculo", "nota_debito")
    ).annotate(
        encargo_aplicado=Subquery(encargos.annotate(t=Sum("valor")).values("t")),
        encargos_ativos=Subquery(encargos.annotate(n=Count("pk")).values("n")),
        # o aviso de cancelar a ND só fala de multas quando ela tem multa
        nd_com_multas=Exists(
            ItemNotaDebito.objects.filter(
                nota_debito_id=OuterRef("nota_debito_id"), multa__isnull=False
            )
        ),
    )
    status = request.GET.get("status", "abertas")
    origem = request.GET.get("origem", "")
    if status == "abertas":
        # judicial também é dívida em aberto: aparece, com o selo "Na Justiça"
        qs = qs.filter(status__in=Cobranca.STATUS_DEVIDOS)
    elif status != "todas":
        qs = qs.filter(status=status)
    if origem:
        qs = qs.filter(origem=origem)
    # Em aberto: mais antigas primeiro (cobrar o atraso mais velho); histórico
    # ("todas"/"pago"/"cancelada"): mais recentes primeiro — senão as cobranças
    # da semana ficariam escondidas na última página.
    cronologica = status in ("abertas", "pendente", "parcial", "atrasado", "judicial")
    qs = qs.order_by(*(("vencimento", "pk") if cronologica else ("-vencimento", "-pk")))
    pagina = Paginator(qs, COBRANCAS_POR_PAGINA).get_page(request.GET.get("pagina"))
    hoje = timezone.localdate()
    abertas_por_cliente = _abertas_por_cliente({c.cliente_id for c in pagina.object_list})
    clientes_com_link = set()
    linhas = []
    for cobranca in pagina.object_list:
        # só aluguel; com 5% já aplicado, sugere o complemento até 10% (D31)
        encargo = (
            services.sugerir_encargo(
                cobranca,
                hoje,
                ja_aplicado=cobranca.encargo_aplicado or ZERO,
                ativos=cobranca.encargos_ativos or 0,
            )
            if cobranca.status == Cobranca.Status.ATRASADO
            else None
        )
        # judicial/quitada saem vazias (whatsapp._cobravel)
        cobranca.link_whatsapp = whatsapp.link_cobranca(cobranca, hoje)
        # "cobrar tudo": uma mensagem com todos os itens, na 1ª linha cobrável do
        # cliente — nunca numa linha judicial/paga (lá o aviso é "sem WhatsApp")
        abertas = abertas_por_cliente.get(cobranca.cliente_id, [])
        cobranca.link_whatsapp_cliente = ""
        cobranca.qtd_abertas_cliente = len(abertas)
        if (
            cobranca.link_whatsapp
            and len(abertas) > 1
            and cobranca.cliente_id not in clientes_com_link
        ):
            clientes_com_link.add(cobranca.cliente_id)
            cobranca.link_whatsapp_cliente = whatsapp.link_cliente(cobranca.cliente, abertas, hoje)
        linhas.append((cobranca, encargo))
    return render(
        request,
        "financeiro/cobrancas.html",
        {
            "linhas": linhas,
            "pagina": pagina,
            "filtros": urlencode({"status": status, "origem": origem}),
            # as ações das linhas voltam para estes filtros e esta página (C29)
            "voltar": urlencode({"status": status, "origem": origem, "pagina": pagina.number}),
            "ordem_rotulo": "mais antigas primeiro" if cronologica else "mais recentes primeiro",
            "status": status,
            "origem": origem,
            "origens": Cobranca.Origem.choices,
            "statuses": Cobranca.Status.choices,
            "origens_cancelaveis": ORIGENS_CANCELAVEIS,
        },
    )


def aplicar_encargo(request, cobranca_id):
    cobranca = get_object_or_404(Cobranca, pk=cobranca_id)
    if request.method == "POST":
        try:
            # aceita "5%"/"10%": vira reais sobre o que falta pagar do aluguel
            valor = services.valor_do_encargo(request.POST.get("valor", ""), cobranca)
            if valor is None:
                valor = _dinheiro(
                    request.POST.get("valor", ""),
                    "Valor do encargo inválido — digite só números (ex.: 45,00) ou 5% / 10%.",
                )
            services.aplicar_encargo(cobranca, valor)
            messages.success(request, f"Encargo de {brl(valor)} aplicado.")
        except (ValidationError, KeyError, InvalidOperation) as erro:
            _erro_no_painel(
                request,
                "; ".join(
                    getattr(
                        erro,
                        "messages",
                        ["Valor do encargo inválido — digite só números (ex.: 45,00)."],
                    )
                ),
            )
    return _voltar_ao_painel(request, cobranca)


def marcar_judicial(request, cobranca_id):
    cobranca = get_object_or_404(Cobranca, pk=cobranca_id)
    if request.method == "POST":
        if cobranca.status not in Cobranca.STATUS_EM_ABERTO:
            # paga/cancelada/já judicial: não reabrir dívida encerrada por clique repetido
            _erro_no_painel(request, "Só cobrança em aberto pode ir para cobrança judicial.")
            return _voltar_ao_painel(request, cobranca)
        cobranca.status = Cobranca.Status.JUDICIAL
        cobranca.save(update_fields=["status"])
        messages.success(request, f"'{cobranca.descricao}' marcada como cobrança judicial.")
    return _voltar_ao_painel(request, cobranca)


def reabrir_judicial(request, cobranca_id):
    """Volta uma cobrança judicial ao fluxo normal — clique errado tem volta."""
    cobranca = get_object_or_404(Cobranca, pk=cobranca_id)
    if request.method == "POST":
        if cobranca.status != Cobranca.Status.JUDICIAL:
            _erro_no_painel(request, "Esta cobrança não está em cobrança judicial.")
        else:
            cobranca.status = Cobranca.Status.PENDENTE
            cobranca.save(update_fields=["status"])
            cobranca.atualizar_status()  # recalcula: pago/parcial/atrasado/pendente
            services.atualizar_inadimplencia(cobranca.cliente)
            messages.success(
                request,
                f"'{cobranca.descricao}' saiu da cobrança judicial "
                f"(status: {cobranca.get_status_display()}).",
            )
    return _voltar_ao_painel(request, cobranca)


def cancelar(request, cobranca_id):
    """Cancela uma cobrança automática que o dono decidiu não cobrar (ex.: excedente de km)."""
    cobranca = get_object_or_404(Cobranca, pk=cobranca_id)
    if request.method == "POST":
        if cobranca.origem not in ORIGENS_CANCELAVEIS:
            _erro_no_painel(
                request,
                "Só cobranças de excedente, encargo, nota de débito ou avulsas podem ser "
                "canceladas.",
            )
        elif cobranca.status == Cobranca.Status.JUDICIAL:
            _erro_no_painel(request, "Cobrança em cobrança judicial não pode ser cancelada.")
        elif cobranca.total_quitado > ZERO:
            _erro_no_painel(request, "Cobrança com pagamento aplicado não pode ser cancelada.")
        elif cobranca.status == Cobranca.Status.CANCELADA:
            messages.info(request, f"Esta cobrança já estava cancelada: {cobranca.descricao}.")
        elif cobranca.nota_debito_id:
            _cancelar_nota_debito(request, cobranca)
        else:
            cobranca.status = Cobranca.Status.CANCELADA
            cobranca.save(update_fields=["status"])
            messages.success(request, f"Cancelada (não será cobrada): {cobranca.descricao}.")
    return _voltar_ao_painel(request, cobranca)


@transaction.atomic
def _cancelar_nota_debito(request, cobranca):
    """Decisão D05: cancela a ND inteira e solta as multas dela.

    O número da ND fica usado (a nota aparece como Cancelada) e as multas voltam
    a ficar "A cobrar" — com o resultado do recurso livre para ser registrado.
    O dono emite uma ND nova só com o que continua valendo.
    """
    cobranca.status = Cobranca.Status.CANCELADA
    cobranca.save(update_fields=["status"])
    soltas = ItemNotaDebito.objects.filter(
        nota_debito_id=cobranca.nota_debito_id, multa__isnull=False
    ).update(multa=None)
    nd = cobranca.nota_debito
    texto = f"ND {nd.numero:03d} cancelada — não será cobrada do cliente."
    if soltas:
        texto += (
            f" {soltas} multa(s) voltaram para 'A cobrar': registre o resultado do recurso "
            "e, se ainda houver o que cobrar, emita uma ND nova."
        )
    # encargo por atraso de ND cancelada some junto, se ainda não teve pagamento
    # nem foi para a cobrança judicial (pendência D05 da etapa 15)
    encargos = Cobranca.com_quitacao_anotada(
        cobranca.encargos.filter(status__in=Cobranca.STATUS_EM_ABERTO)
    )
    sem_pagamento = [e.pk for e in encargos if e.total_quitado <= ZERO]
    if sem_pagamento:
        Cobranca.objects.filter(pk__in=sem_pagamento).update(status=Cobranca.Status.CANCELADA)
        texto += (
            " O encargo por atraso dela também foi cancelado."
            if len(sem_pagamento) == 1
            else f" Os {len(sem_pagamento)} encargos por atraso dela também foram cancelados."
        )
    services.atualizar_inadimplencia(cobranca.cliente)
    messages.success(request, texto)
    if cobranca.encargos.filter(status__in=Cobranca.STATUS_DEVIDOS).exists():
        messages.warning(
            request,
            "Esta ND tem encargo por atraso que já recebeu parte do pagamento ou está na "
            "cobrança judicial — ele continua como está. Resolva esse encargo à parte.",
        )


def _id_valido(texto):
    """pk vinda da URL: só algarismos ASCII e de tamanho sensato.

    str.isdigit() aceita "²" e números gigantes estouram o inteiro do banco —
    os dois virariam erro 500 em vez de serem ignorados.
    """
    return bool(texto) and texto.isascii() and texto.isdigit() and len(texto) <= 9


def _resumo_do_recibo(request, cliente):
    """?recibo=<pk>: o aviso pós-lançamento (o que quitou, sobra, recibo, desfazer)."""
    recibo_id = request.GET.get("recibo", "")
    if not (cliente and _id_valido(recibo_id)):
        return None
    recebimento = (
        services.anotar_lancamento(Recebimento.objects.filter(pk=recibo_id, cliente=cliente))
        .select_related("cliente")
        .first()
    )
    if recebimento is None:  # desfeito ou de outro cliente: some sem erro
        return None
    resumo = services.resumo_recebimento(recebimento)
    # crédito usado junto com o dinheiro novo (?recibo_credito=<pk>): entra no resumo
    credito_id = request.GET.get("recibo_credito", "")
    junto = (
        Recebimento.objects.filter(
            pk=credito_id, cliente=cliente, forma=Recebimento.Forma.CREDITO
        ).first()
        if _id_valido(credito_id)
        else None
    )
    resumo["recebimento_credito"] = junto
    if junto is not None:
        services.juntar_resumo_do_credito(resumo, junto)
    resumo["link_whatsapp"] = recibo.link_recibo(resumo)
    resumo["pode_desfazer"] = services.pode_desfazer(recebimento)
    return resumo


def _opcoes_de_clientes():
    """Seletor de cliente do Receber e da ND: "Nome — placa — deve R$ X".

    Homônimos ("Lucas" e "Lucas Martins Ferreira") se distinguem pela placa e
    pelo saldo; quem tem valor em aberto vem primeiro. Três consultas, sem
    importar quantos clientes.
    Devolve [(rótulo do grupo, [(pk, rótulo da opção), ...]), ...].
    """
    from apps.alocacoes.models import Alocacao

    placas = {}
    for cliente_id, placa in (
        Alocacao.objects.filter(status=Alocacao.Status.ATIVA)
        .order_by("data_inicio")
        .values_list("cliente_id", "veiculo__placa")
    ):
        placas.setdefault(cliente_id, []).append(placa)
    # mesma conta do "Quem deve" e do Painel: "deve hoje" é só o vencido;
    # o que ainda vai vencer aparece à parte (antes o seletor somava os dois)
    hoje = timezone.localdate()
    vencido, a_vencer = {}, {}
    for cobranca in Cobranca.com_quitacao_anotada(
        Cobranca.objects.filter(status__in=Cobranca.STATUS_DEVIDOS)
    ):
        if cobranca.saldo > ZERO:
            soma = vencido if cobranca.vencimento <= hoje else a_vencer
            soma[cobranca.cliente_id] = soma.get(cobranca.cliente_id, ZERO) + cobranca.saldo
    com_valor, demais = [], []
    for pk, nome in Cliente.objects.order_by("nome").values_list("pk", "nome"):
        partes = [nome, *placas.get(pk, [])]
        if pk in vencido or pk in a_vencer:
            partes.append(texto_do_devido(vencido.get(pk, ZERO), a_vencer.get(pk, ZERO)))
            com_valor.append((pk, " — ".join(partes)))
        else:
            demais.append((pk, " — ".join(partes)))
    grupos = []
    if com_valor:
        grupos.append(("Com valor em aberto", com_valor))
    if demais:
        grupos.append(("Demais clientes" if com_valor else "Clientes", demais))
    return grupos


def texto_do_devido(vencido, a_vencer):
    """O mesmo texto em todas as telas: "deve hoje R$ X (+ R$ Y a vencer)"."""
    if vencido > ZERO and a_vencer > ZERO:
        return f"deve hoje {brl(vencido)} (+ {brl(a_vencer)} a vencer)"
    if vencido > ZERO:
        return f"deve hoje {brl(vencido)}"
    return f"{brl(a_vencer)} a vencer"


def _vencido_e_a_vencer(cobrancas, hoje):
    """(vencido até hoje, a vencer) das cobranças devidas — mesma conta do "Quem deve"."""
    vencido = sum((c.saldo for c in cobrancas if c.vencimento <= hoje), ZERO)
    a_vencer = sum((c.saldo for c in cobrancas if c.vencimento > hoje), ZERO)
    return vencido, a_vencer


def _ler_aplicacoes(post, cliente, cobrancas_abertas):
    """[(cobrança, valor)] de TODAS as chaves aplicar_<pk> do formulário.

    Antes só as cobranças recarregadas eram lidas: um valor digitado numa
    cobrança que outra aba acabou de quitar era descartado em silêncio e o
    dinheiro virava crédito. Agora ele trava o lançamento com aviso.
    """
    por_pk = {str(c.pk): c for c in cobrancas_abertas}
    aplicacoes = []
    for chave in post:
        if not chave.startswith("aplicar_") or chave.startswith("aplicar_encargo_"):
            continue
        pk = chave.removeprefix("aplicar_")
        bruto = post.get(chave, "").strip()
        if not bruto or not _id_valido(pk):
            continue
        cobranca = por_pk.get(pk)
        descricao = cobranca.descricao if cobranca else None
        if cobranca is None:
            antiga = Cobranca.objects.filter(pk=pk, cliente=cliente).first()
            descricao = antiga.descricao if antiga else "Uma das cobranças"
        valor = _dinheiro(bruto, f"Valor a aplicar inválido em '{descricao}' (ex.: 650,00).")
        if valor <= ZERO:
            continue
        if cobranca is None or cobranca.saldo <= ZERO:
            raise ValidationError(
                f"'{descricao}' já não está em aberto (foi quitada ou cancelada). "
                "A tela foi atualizada — confira e lance de novo."
            )
        aplicacoes.append((cobranca, valor))
    return aplicacoes


def _encargos_sugeridos(cobrancas_abertas, hoje):
    """{pk do aluguel atrasado: encargo sugerido} — para cobrar o atraso junto no Receber."""
    sugeridos = {}
    for cobranca in cobrancas_abertas:
        if cobranca.status != Cobranca.Status.ATRASADO:
            continue
        valor = services.sugerir_encargo(
            cobranca,
            hoje,
            ja_aplicado=getattr(cobranca, "encargo_aplicado", None) or ZERO,
            ativos=getattr(cobranca, "encargos_ativos", None) or 0,
        )
        if valor > ZERO:
            sugeridos[cobranca.pk] = valor
    return sugeridos


def _criar_encargos_do_post(post, cobrancas_abertas, sugeridos, hoje):
    """Encargos marcados em "incluir" no Receber: cria cada um e devolve as aplicações.

    Roda dentro da mesma transação do recebimento: se o lançamento falhar,
    o encargo também não fica. Devolve (aplicações, encargos criados).
    """
    aplicacoes, criados = [], []
    por_pk = {c.pk: c for c in cobrancas_abertas}
    for pk in sugeridos:
        if post.get(f"incluir_encargo_{pk}") != "1":
            continue
        aluguel = por_pk[pk]
        valor = _dinheiro(
            post.get(f"encargo_{pk}", ""),
            f"Valor do encargo inválido em '{aluguel.descricao}' (ex.: 65,00).",
        )
        if valor <= ZERO:
            continue
        encargo = services.aplicar_encargo(aluguel, valor, hoje)
        bruto = post.get(f"aplicar_encargo_{pk}", "").strip()
        aplicar = (
            _dinheiro(bruto, f"Valor a aplicar inválido no encargo de '{aluguel.descricao}'.")
            if bruto
            else ZERO
        )
        if aplicar > valor:
            raise ValidationError(
                f"Aplicação de {brl(aplicar)} acima do encargo ({brl(valor)}) de "
                f"'{aluguel.descricao}'."
            )
        if aplicar > ZERO:
            aplicacoes.append((encargo, aplicar))
        criados.append(encargo)
    return aplicacoes, criados


def receber(request):
    """Baixa de recebimento com travas de lançamento (docs.md §4.3).

    Cobrança judicial também aparece aqui (STATUS_DEVIDOS): acordo pago
    precisa de onde ser lançado. No erro, a tela volta com tudo que foi
    digitado e diz qual campo falhou — nada de redigitar a distribuição.

    Atalhos (botões "Recebi"): ?cliente=<pk> já escolhe o cliente,
    ?cobranca=<pk> já preenche o valor com o que falta daquela cobrança e
    ?usar_credito=1 já liga o pagamento com o crédito guardado ("Quem deve").
    Depois de lançar, ?recibo=<pk> mostra o resumo, o recibo e o desfazer.
    """
    hoje = timezone.localdate()
    cliente = None
    cobrancas_abertas = []
    cliente_id = request.GET.get("cliente") or request.POST.get("cliente")
    if cliente_id:
        if not _id_valido(cliente_id):
            raise Http404("Cliente inválido.")
        cliente = get_object_or_404(Cliente, pk=cliente_id)
        cobrancas_abertas = _cobrancas_devidas(cliente)
    post = None
    if request.method == "POST" and cliente:
        post = request.POST
        try:
            valor = _dinheiro(
                post.get("valor", ""),
                "Valor recebido inválido — digite só números, com vírgula "
                "para os centavos (ex.: 1250,00).",
            )
            data = ler_data(post.get("data", ""), "Data do recebimento")
            forma = post.get("forma", "pix")
            usar_credito = ZERO
            if post.get("usar_credito_junto") == "1" and forma != Recebimento.Forma.CREDITO:
                usar_credito = max(MovimentoCredito.saldo_do_cliente(cliente), ZERO)
            # toque duplo: o 2º envio acha a cobrança já paga — o aviso certo é o de repetido
            repetido = services._recebimento_recente(cliente, data, valor)
            if repetido:
                raise services.RecebimentoRepetido(repetido)
            aplicacoes = _ler_aplicacoes(post, cliente, cobrancas_abertas)
            with transaction.atomic():
                nos_encargos, encargos = _criar_encargos_do_post(
                    post, cobrancas_abertas, _encargos_sugeridos(cobrancas_abertas, hoje), hoje
                )
                aplicacoes += nos_encargos
                recebimento = services.registrar_recebimento(
                    cliente=cliente,
                    data=data,
                    valor=valor,
                    forma=forma,
                    aplicacoes=aplicacoes,
                    sobra_destino=post.get("sobra_destino", "credito"),
                    observacoes=post.get("observacoes", ""),
                    bloquear_repetido=True,
                    usar_credito=usar_credito,
                    mes_passado_confirmado=post.get("confirmar_mes_passado") == "1",
                    hoje=hoje,
                )
                # o "desfazer" deste lançamento também cancela o encargo criado aqui
                Cobranca.objects.filter(pk__in=[e.pk for e in encargos]).update(
                    observacoes=services.marca_encargo(recebimento)
                )
            if recebimento.forma == Recebimento.Forma.CREDITO:
                messages.success(request, f"Crédito de {brl(valor)} usado para {cliente.nome}.")
            else:
                messages.success(
                    request, f"Recebimento de {brl(valor)} lançado para {cliente.nome}."
                )
            destino = f"{request.path}?cliente={cliente.pk}&recibo={recebimento.pk}"
            if recebimento.recebimento_credito:
                destino += f"&recibo_credito={recebimento.recebimento_credito.pk}"
            # aberto do acerto da caução (?voltar=): o recibo mostra o botão de volta
            if _voltar_seguro(request)[0]:
                destino += "&" + urlencode({"voltar": _voltar_seguro(request)[0]})
            return redirect(destino)
        except services.RecebimentoRepetido as repetido:
            # toque duplo: mostra o resumo do lançamento que já valeu
            messages.warning(request, repetido.messages[0])
            return redirect(f"{request.path}?cliente={cliente.pk}&recibo={repetido.recebimento.pk}")
        except ValidationError as erro:
            messages.error(request, "; ".join(erro.messages))
            # a transação foi desfeita, mas os objetos em memória podem ter status
            # de uma aplicação que não valeu: a tela volta com o que está no banco
            cobrancas_abertas = _cobrancas_devidas(cliente)
    elif cliente and _id_valido(request.GET.get("cobranca", "")):
        # "Recebi" de uma cobrança: valor = o que falta dela, já aplicado nela
        escolhida = next(
            (c for c in cobrancas_abertas if str(c.pk) == request.GET["cobranca"]), None
        )
        if escolhida and escolhida.saldo > ZERO:
            saldo = exibir_dinheiro(escolhida.saldo)
            post = {"valor": saldo, f"aplicar_{escolhida.pk}": saldo}
    credito = MovimentoCredito.saldo_do_cliente(cliente) if cliente else ZERO
    sugeridos = _encargos_sugeridos(cobrancas_abertas, hoje)
    linhas = []
    for cobranca in cobrancas_abertas:
        # encargo por atraso que dá para cobrar junto (aluguel atrasado)
        cobranca.encargo_sugerido = sugeridos.get(cobranca.pk)
        cobranca.encargo_digitado = (post.get(f"encargo_{cobranca.pk}", "") if post else "") or (
            exibir_dinheiro(sugeridos[cobranca.pk]) if cobranca.pk in sugeridos else ""
        )
        cobranca.encargo_incluido = bool(post and post.get(f"incluir_encargo_{cobranca.pk}") == "1")
        cobranca.encargo_para_aplicar = (
            post.get(f"aplicar_encargo_{cobranca.pk}", "") if post else ""
        )
        linhas.append((cobranca, post.get(f"aplicar_{cobranca.pk}", "") if post else ""))
    vencido, a_vencer = _vencido_e_a_vencer(cobrancas_abertas, hoje)
    # "Vencido" só é o que venceu ANTES de hoje; o que vence hoje ainda não atrasou
    atrasado = sum((c.saldo for c in cobrancas_abertas if c.vencimento < hoje), ZERO)
    resumo = _resumo_do_recibo(request, cliente)
    return render(
        request,
        "financeiro/receber.html",
        {
            "grupos_clientes": _opcoes_de_clientes(),
            "cliente": cliente,
            "linhas": linhas,
            "credito": credito,
            "devido": vencido + a_vencer,
            "vencido": vencido,
            "a_vencer": a_vencer,
            "hoje": hoje,
            # data antes do dia 1º deste mês mexe num DAS anterior: pede confirmação
            "inicio_do_mes": services.inicio_do_mes(hoje),
            "post": post,
            # crédito não é forma de pagamento: tem botão próprio na tela
            "formas": [f for f in Recebimento.Forma.choices if f[0] != Recebimento.Forma.CREDITO],
            # reforço de caução só para quem tem caução de contrato ativo
            "tem_caucao": bool(cliente and services._caucao_ativa(cliente) is not None),
            "abrir_com_credito": request.GET.get("usar_credito") == "1" and not post,
            "resumo": resumo,
            # o cartão verde do recibo já diz "Pagamento lançado": sem o aviso repetido
            "sucesso_no_cartao": bool(resumo),
            "atrasado": atrasado,
            "voltar": _voltar_seguro(request),
        },
    )


def desfazer_recebimento(request, recebimento_id):
    """Botão "Desfazer este lançamento" — só no mesmo dia (services.pode_desfazer)."""
    recebimento = Recebimento.objects.select_related("cliente").filter(pk=recebimento_id).first()
    if recebimento is None:
        # toque duplo no Desfazer (ou outra aba): o 1º já apagou — avisa em vez de "página
        # não encontrada"; o histórico ainda sabe de qual cliente era
        apagado = Recebimento.history.filter(id=recebimento_id).order_by("-history_date").first()
        if apagado is None:
            raise Http404("Lançamento não encontrado.")
        messages.info(request, "Este lançamento já foi desfeito.")
        return redirect(f"{reverse('financeiro:receber')}?cliente={apagado.cliente_id}")
    cliente = recebimento.cliente
    if request.method != "POST":
        return redirect(f"{reverse('financeiro:receber')}?cliente={cliente.pk}")
    valor = recebimento.valor
    # pagamento com dinheiro + crédito: os dois lançamentos se desfazem juntos
    junto_id = request.POST.get("junto", "")
    junto = (
        Recebimento.objects.filter(
            pk=junto_id, cliente=cliente, forma=Recebimento.Forma.CREDITO
        ).first()
        if _id_valido(junto_id)
        else None
    )
    # desfazer pelo extrato volta para o extrato — no erro, sem o cartão verde
    # "Pagamento lançado" do Receber, que fazia parecer que tinha dado certo
    do_extrato = request.POST.get("voltar") == "extrato"
    extrato = reverse("financeiro:extrato_cliente", args=[cliente.pk])
    try:
        with transaction.atomic():
            # o crédito sai primeiro: se ele pagou o encargo cobrado junto, o encargo
            # fica sem pagamento e o desfazer do dinheiro consegue cancelá-lo
            if junto is not None:
                services.desfazer_recebimento(junto)
                valor += junto.valor
            services.desfazer_recebimento(recebimento)
    except ValidationError as erro:
        texto = "; ".join(erro.messages)
        if "já foi usado" in texto and "crédito" in texto:
            texto += (
                " No extrato, ele aparece como “Abatimento com crédito do cliente”, "
                "com o “↩ desfazer” ao lado (se foi lançado hoje)."
            )
        messages.error(request, texto)
        if do_extrato:
            return redirect(extrato)
        destino = f"{reverse('financeiro:receber')}?cliente={cliente.pk}&recibo={recebimento.pk}"
        if junto is not None:
            destino += f"&recibo_credito={junto.pk}"
        return redirect(destino)
    messages.success(
        request,
        f"Lançamento de {brl(valor)} de {cliente.nome} desfeito — as cobranças voltaram "
        "a ficar em aberto. Se precisar, lance de novo com os dados certos.",
    )
    if do_extrato:
        return redirect(extrato)
    return redirect(f"{reverse('financeiro:receber')}?cliente={cliente.pk}")


#: filtros da lista de NDs: (valor na URL, rótulo, status da cobrança da ND)
FILTROS_ND = [
    ("abertas", "Em aberto", Cobranca.STATUS_DEVIDOS),
    ("pagas", "Pagas", [Cobranca.Status.PAGO]),
    ("canceladas", "Canceladas", [Cobranca.Status.CANCELADA]),
    ("todas", "Todas", None),
]


def nds(request):
    """Lista de NDs — abre nas que ainda estão em aberto (com "✓ Recebi")."""
    situacao = request.GET.get("situacao", "abertas")
    filtro = next((f for f in FILTROS_ND if f[0] == situacao), FILTROS_ND[0])
    situacao = filtro[0]
    # cobrança com o quitado anotado: o botão "Enviar ao motorista" olha o saldo
    # de cada ND sem 2 consultas por linha
    qs = NotaDebito.objects.select_related("cliente").prefetch_related(
        "itens", Prefetch("cobranca", queryset=Cobranca.com_quitacao_anotada())
    )
    if filtro[2] is not None:
        qs = qs.filter(cobranca__status__in=filtro[2])
    notas = list(qs)
    hoje = timezone.localdate()
    # crédito guardado do cliente: entra na mensagem e lembra de usar antes de lançar Pix
    creditos = creditos_por_cliente({nd.cliente_id for nd in notas})
    for nd in notas:
        nd.credito = creditos.get(nd.cliente_id, ZERO)
        nd.link_whatsapp = whatsapp.link_nd(nd, hoje, credito=nd.credito)
        # cobrável mas sem telefone: avisa em vez de sumir com o botão sem explicação
        cobranca = getattr(nd, "cobranca", None)
        nd.sem_telefone = (
            not nd.link_whatsapp
            and cobranca is not None
            and cobranca.status in Cobranca.STATUS_EM_ABERTO
            and cobranca.saldo > ZERO
            and not whatsapp.telefone_whatsapp(nd.cliente.telefone)
        )
        # D05: ND sem pagamento e fora do judicial pode ser cancelada inteira
        nd.pode_cancelar = (
            cobranca is not None
            and cobranca.status in Cobranca.STATUS_EM_ABERTO
            and cobranca.total_quitado == ZERO
        )
        nd.pode_receber = (
            cobranca is not None
            and cobranca.status in Cobranca.STATUS_DEVIDOS
            and cobranca.saldo > ZERO
        )
        # o aviso de cancelar só fala de multas quando a ND tem multa
        nd.tem_multas = any(item.multa_id for item in nd.itens.all())
    return render(
        request,
        "financeiro/nds.html",
        {
            "notas": notas,
            "hoje": hoje,
            "situacao": situacao,
            "filtros": [(valor, rotulo) for valor, rotulo, _ in FILTROS_ND],
        },
    )


def _voltar_seguro(request):
    """?voltar= (ou o campo "voltar"): de onde a tela foi aberta — o acerto da
    caução, a devolução do carro —, para voltar para lá depois de lançar.
    Só caminho deste site ("/financeiro/caucoes/3/"), nunca outro endereço.
    Devolve (endereço, rótulo do botão) ou (None, "")."""
    voltar = request.POST.get("voltar") or request.GET.get("voltar", "")
    if not (
        voltar.startswith("/")
        and not voltar.startswith("//")
        and not voltar.startswith("/\\")
        and url_has_allowed_host_and_scheme(
            voltar, allowed_hosts={request.get_host()}, require_https=request.is_secure()
        )
    ):
        return None, ""
    if voltar.startswith(reverse("financeiro:caucoes")):
        return voltar, "Voltar ao acerto da caução"
    if "/encerrar/" in voltar:
        return voltar, "Voltar à devolução do carro"
    return voltar, "Voltar para a tela anterior"


def nd_nova(request):
    """Emissão de ND idempotente: token único por formulário barra o clique duplo,
    e o erro devolve a tela com os itens digitados preservados.

    ?cliente=<pk> já escolhe o cliente; ?voltar=<caminho> volta para lá depois
    de emitir (avaria achada na devolução: acerto da caução ou encerrar)."""
    post = None
    if request.method == "POST":
        post = request.POST
        token = post.get("token", "")
        try:
            with _envio_unico(request, "nd_tokens_usados", token) as repetido:
                if repetido:
                    # segunda submissão do mesmo formulário: a ND já foi emitida
                    messages.warning(
                        request,
                        "Esta nota de débito já tinha sido emitida — o clique duplo foi "
                        "ignorado e nada foi cobrado em dobro.",
                    )
                    return redirect("financeiro:nds")
                cliente = get_object_or_404(Cliente, pk=post["cliente"])
                descricoes = post.getlist("item_descricao")
                valores = post.getlist("item_valor")
                itens = []
                for numero, (descricao, valor) in enumerate(
                    zip(descricoes, valores, strict=False), start=1
                ):
                    descricao = services.limpar_texto(descricao, maximo=10_000)
                    if not descricao and not valor.strip():
                        continue  # linha totalmente vazia: só some
                    # linha pela metade não pode sumir em silêncio: a tela e a
                    # confirmação mostravam um total e a ND saía com outro
                    if not descricao:
                        valor_lido = _dinheiro(valor, f"Valor inválido no item {numero}.")
                        raise ValidationError(
                            f"Item {numero} tem valor ({brl(valor_lido)}) mas não tem "
                            "descrição: escreva o que é, ou apague o valor."
                        )
                    if not valor.strip():
                        raise ValidationError(f'"{descricao[:40]}" está sem valor.')
                    if len(descricao) > MAX_DESCRICAO_ITEM_ND:
                        raise ValidationError(
                            f'A descrição do item "{descricao[:40]}…" tem {len(descricao)} '
                            f"letras; o máximo é {MAX_DESCRICAO_ITEM_ND} letras. Encurte o texto."
                        )
                    valor_item = _dinheiro(
                        valor, f'Valor inválido no item "{descricao[:40]}" (ex.: 150,00).'
                    )
                    if valor_item <= ZERO:
                        raise ValidationError(f'"{descricao[:40]}" está sem valor.')
                    itens.append((descricao, valor_item))
                try:
                    data_emissao = ler_data(post.get("data_emissao"), "Data de emissão")
                    vencimento = (
                        ler_data(post["vencimento"], "Data de vencimento")
                        if post.get("vencimento", "").strip()
                        else None
                    )
                except ValidationError as erro:
                    raise ValidationError("Não deu para emitir: " + erro.messages[0]) from None
                nd = services.emitir_nota_debito(
                    cliente=cliente,
                    data_emissao=data_emissao,
                    vencimento=vencimento,
                    itens=itens,
                    observacoes=post.get("observacoes", ""),
                )
            messages.success(request, f"ND {nd.numero:03d} emitida ({brl(nd.total)}).")
            return redirect(_voltar_seguro(request)[0] or "financeiro:nds")
        except ValidationError as erro:
            messages.error(request, "; ".join(erro.messages))
        except (KeyError, InvalidOperation, ValueError):
            messages.error(
                request,
                "Não deu para emitir: escolha o cliente, confira as datas de emissão e de "
                "vencimento e "
                "preencha ao menos um item com descrição e valor (ex.: 150,00).",
            )
    itens_iniciais = [{"d": "", "v": ""}]
    if post:
        digitados = [
            {"d": descricao, "v": valor}
            for descricao, valor in zip(
                post.getlist("item_descricao"), post.getlist("item_valor"), strict=False
            )
        ]
        itens_iniciais = digitados or itens_iniciais
    # ?cliente=<pk>: link da ficha, da caução ou da devolução já escolhe o cliente
    cliente_inicial = post.get("cliente", "") if post else request.GET.get("cliente", "")
    if not _id_valido(cliente_inicial) or not Cliente.objects.filter(pk=cliente_inicial).exists():
        cliente_inicial = ""
    return render(
        request,
        "financeiro/nd_nova.html",
        {
            "cliente_inicial": str(cliente_inicial),
            "voltar": _voltar_seguro(request),
            "grupos_clientes": _opcoes_de_clientes(),
            # o formulário calcula o vencimento sugerido ao trocar cliente/emissão
            "dias_pagamento": {
                str(pk): dia for pk, dia in services.dias_pagamento_por_cliente().items()
            },
            "hoje": timezone.localdate(),
            "post": post,
            "itens_iniciais": itens_iniciais,
            "token": uuid4().hex,
        },
    )


def caucoes(request):
    lista = Caucao.com_saldos_anotados(
        Caucao.objects.select_related("alocacao__cliente", "alocacao__veiculo")
    )
    return render(request, "financeiro/caucoes.html", {"caucoes": lista})


def caucao_detalhe(request, caucao_id):
    caucao = get_object_or_404(
        Caucao.com_saldos_anotados(
            Caucao.objects.select_related("alocacao__cliente", "alocacao__veiculo")
        ),
        pk=caucao_id,
    )
    cliente = caucao.alocacao.cliente
    token = request.POST.get("token", "")
    post = None
    if request.method == "POST":
        post = request.POST
        try:
            with _envio_unico(request, "caucao_tokens_usados", token) as repetido:
                if repetido:
                    # segundo envio do mesmo formulário (clique duplo): já foi gravado (C06)
                    messages.warning(
                        request,
                        "Esta movimentação já tinha sido registrada — o clique duplo foi "
                        "ignorado e nada foi lançado em dobro.",
                    )
                    return redirect("financeiro:caucao_detalhe", caucao.pk)
                tipo = request.POST["tipo"]
                if tipo not in MovimentacaoCaucao.Tipo.values:
                    raise ValidationError('Escolha em "O que aconteceu?" o tipo da movimentação.')
                valor = _dinheiro(
                    request.POST.get("valor", ""),
                    "Valor inválido — digite só números, com vírgula nos centavos (ex.: 1.500,00).",
                )
                if valor <= ZERO:
                    raise ValidationError("O valor deve ser maior que zero.")
                data = ler_data(request.POST.get("data", ""))
                observacoes = request.POST.get("observacoes", "")
                if tipo == MovimentacaoCaucao.Tipo.DESCONTO:
                    cobranca_id = request.POST.get("cobranca", "")
                    cobranca = (
                        Cobranca.objects.filter(pk=cobranca_id, cliente=cliente).first()
                        if _id_valido(cobranca_id)
                        else None
                    )
                    if cobranca is None:
                        raise ValidationError("Escolha a cobrança que a caução vai pagar.")
                    services.descontar_da_caucao(
                        caucao,
                        cobranca,
                        valor,
                        data,
                        observacoes=observacoes,
                        mes_passado_confirmado=request.POST.get("confirmar_mes_passado") == "1",
                    )
                else:
                    services.registrar_movimentacao_caucao(
                        caucao,
                        tipo,
                        valor,
                        data,
                        forma=request.POST.get("forma", ""),
                        observacoes=observacoes,
                        acima_do_combinado_confirmado=request.POST.get("confirmar_acima") == "1",
                        devolver_mesmo_assim=request.POST.get("devolver_mesmo_assim") == "1",
                    )
            messages.success(
                request, f"Movimentação de {brl(valor)} registrada na caução de {cliente.nome}."
            )
            return redirect("financeiro:caucao_detalhe", caucao.pk)
        except ValidationError as erro:
            messages.error(request, "; ".join(erro.messages))
        except (KeyError, InvalidOperation, ValueError):
            messages.error(
                request,
                "Não deu para registrar: confira o valor (ex.: 1.500,00), a data e, "
                "se for desconto, escolha a cobrança a quitar.",
            )
    hoje = timezone.localdate()
    movimentacoes = list(
        services.anotar_lancamento_caucao(caucao.movimentacoes.select_related("cobranca"))
    )
    for movimentacao in movimentacoes:
        movimentacao.pode_desfazer = services.pode_desfazer_movimentacao(movimentacao, hoje)
    combinado = caucao.alocacao.caucao_valor or ZERO
    deve, ativa = services.pendencias_para_devolver(caucao)
    # "falta depositar" = combinado − o que o cliente já depositou (depósito +
    # reforço), não − saldo: depois de devolver ou descontar, o saldo cai e a
    # tela sugeria um depósito que não aconteceu. Contrato encerrado: nada a depositar.
    falta = max(combinado - caucao.recebido, ZERO) if ativa else ZERO
    pendentes = _cobrancas_devidas(cliente)
    # Carro já devolvido e caução com saldo (ex.: logo depois de encerrar a
    # alocação): a tela já abre no acerto — descontar a dívida ou devolver —,
    # e não em "Cliente depositou a caução".
    tipo_inicial = MovimentacaoCaucao.Tipo.RECEBIMENTO
    valor_inicial = exibir_dinheiro(falta) if falta > ZERO else ""
    if not ativa and caucao.saldo > ZERO:
        com_saldo = [c for c in pendentes if c.saldo > ZERO]
        if com_saldo:
            tipo_inicial = MovimentacaoCaucao.Tipo.DESCONTO
            valor_inicial = exibir_dinheiro(min(caucao.saldo, com_saldo[0].saldo))
        else:
            tipo_inicial = MovimentacaoCaucao.Tipo.DEVOLUCAO
            valor_inicial = exibir_dinheiro(caucao.saldo)
    return render(
        request,
        "financeiro/caucao_detalhe.html",
        {
            "caucao": caucao,
            "cliente": cliente,
            "movimentacoes": movimentacoes,
            # judicial entra: acordo pode ser quitado descontando da caução retida
            "pendentes": pendentes,
            # crédito (pagamento adiantado / semanas não usadas) é do cliente, como a
            # caução: o dono vê os dois juntos antes de devolver ou descontar
            "credito": MovimentoCredito.saldo_do_cliente(cliente),
            "combinado": combinado,
            "falta": falta,
            # o campo Valor já vem com o que falta depositar (o caso mais comum)
            "valor_sugerido": exibir_dinheiro(falta) if falta > ZERO else "",
            "tipo_inicial": (post.get("tipo") if post else "") or tipo_inicial,
            "valor_inicial": valor_inicial,
            "deve": deve,
            "alocacao_ativa": ativa,
            "hoje": hoje,
            "data_minima": caucao.alocacao.data_inicio,
            # desconto com data de mês anterior muda o DAS daquele mês: pede confirmação
            "inicio_do_mes": services.inicio_do_mes(hoje),
            "data_minima_deposito": caucao.alocacao.data_inicio
            - timedelta(days=services.DIAS_DEPOSITO_ANTES_DO_CONTRATO),
            "post": post,
            "token": uuid4().hex,
        },
    )


def desfazer_movimentacao_caucao(request, movimentacao_id):
    """ "↩ desfazer" de uma movimentação de caução lançada hoje (valor ou tipo errado)."""
    movimentacao = (
        MovimentacaoCaucao.objects.select_related("caucao").filter(pk=movimentacao_id).first()
    )
    if movimentacao is None:
        # desfazer em duas abas: o 1º já apagou — avisa em vez de "página não
        # encontrada"; o histórico ainda sabe de qual caução era
        apagada = (
            MovimentacaoCaucao.history.filter(id=movimentacao_id).order_by("-history_date").first()
        )
        if apagada is None:
            raise Http404("Movimentação não encontrada.")
        messages.info(request, "Esta movimentação já foi desfeita.")
        return redirect("financeiro:caucao_detalhe", apagada.caucao_id)
    caucao_id = movimentacao.caucao_id
    if request.method == "POST":
        valor = movimentacao.valor
        tipo = movimentacao.get_tipo_display().lower()
        try:
            services.desfazer_movimentacao_caucao(movimentacao)
            messages.success(request, f"Lançamento de {tipo} de {brl(valor)} desfeito.")
        except ValidationError as erro:
            messages.error(request, "; ".join(erro.messages))
    return redirect("financeiro:caucao_detalhe", caucao_id)


COLUNAS_PAGAMENTOS = [
    "Data",
    "Cliente",
    "Placa",
    "Descrição",
    "Forma",
    "Valor (R$)",
    "Entra no DAS?",
]


def _linhas_pagamentos(pagamentos):
    """Linhas da aba "Pagamentos do mês" + o total, que bate com a 1ª aba."""
    linhas = [
        [
            p["data"].strftime("%d/%m/%Y"),
            p["cliente"],
            p["placa"],
            p["descricao"],
            p["forma"],
            p["valor"],
            "Sim" if p["entra_no_das"] else "Não",
        ]
        for p in pagamentos
    ]
    total = sum((p["valor"] for p in pagamentos), ZERO)
    no_das = sum((p["valor"] for p in pagamentos if p["entra_no_das"]), ZERO)
    linhas.append(["", "", "", "Total (locação + pagamentos diversos)", "", total, ""])
    linhas.append(["", "", "", "Dos quais entram na base do DAS", "", no_das, "Sim"])
    return linhas


def linhas_base_das(ano, mes):
    """Linhas da 1ª aba do DAS (totais) — as mesmas em Financeiro › DAS e em
    Relatórios, para a contabilidade receber a mesma planilha das duas telas."""
    r = services.resumo_fiscal(ano, mes)
    linhas = [["Receita de locação (fatura)", r["locacao"], "Sim"]]
    for rotulo, valor in r["diversos"].items():
        linhas.append([f"Pagamentos diversos — {rotulo} (ND)", valor, "Não"])
    linhas.append(["Caução recebida (não é receita — é do cliente)", r["caucao_recebida"], "Não"])
    linhas.append(["Adiantamentos recebidos (crédito do cliente)", r["adiantamentos"], "Não"])
    linhas.append(["Pagos com crédito do cliente (já contados acima)", r["pago_com_credito"], "-"])
    linhas.append(["Total que entrou no caixa", r["entrou_no_caixa"], "-"])
    return linhas


def linhas_pagamentos_das(ano, mes):
    """Linhas da 2ª aba do DAS: cada pagamento que compõe o DAS e, embaixo, o
    dinheiro do mês que fica fora do DAS (sobra que virou crédito e caução
    recebida), para a aba fechar com o "Total que entrou no caixa" da 1ª."""
    linhas = _linhas_pagamentos(services.pagamentos_do_mes(ano, mes))
    fora = services.entradas_fora_do_das(ano, mes)
    if fora:
        linhas.append(["", "", "", "Entrou no mês, mas fora do DAS:", "", "", ""])
        linhas += [
            [
                p["data"].strftime("%d/%m/%Y"),
                p["cliente"],
                p["placa"],
                p["descricao"],
                p["forma"],
                p["valor"],
                "Não",
            ]
            for p in fora
        ]
    return linhas


def _planilha_das(formato, nome, linhas, linhas_pagamentos):
    """Planilha do DAS para a contabilidade, pronta para o Excel em português (C11).

    xlsx: duas abas — "Base do DAS" (totais) e "Pagamentos do mês" (cada
    pagamento, com total que bate com a 1ª aba). csv: só os totais, com ponto
    e vírgula entre colunas, vírgula nos centavos e BOM UTF-8.
    """
    colunas = ["Grupo", "Valor (R$)", "Entra na base do DAS?"]
    if formato == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Font

        livro = Workbook()
        for indice, (titulo, cabecalho, dados) in enumerate(
            [
                ("Base do DAS", colunas, linhas),
                ("Pagamentos do mês", COLUNAS_PAGAMENTOS, linhas_pagamentos),
            ]
        ):
            aba = livro.active if indice == 0 else livro.create_sheet()
            aba.title = titulo
            aba.append(cabecalho)
            for celula in aba[1]:
                celula.font = Font(bold=True)
            for linha in dados:
                aba.append([float(v) if isinstance(v, Decimal) else v for v in linha])
                for celula in aba[aba.max_row]:
                    if isinstance(celula.value, str):
                        celula.data_type = "s"  # texto digitado nunca vira fórmula
                    elif isinstance(celula.value, float):
                        celula.number_format = "#,##0.00"
            for coluna, titulo_coluna in enumerate(cabecalho, 1):
                largura = max(
                    [len(titulo_coluna)] + [len(str(linha[coluna - 1])) for linha in dados]
                )
                letra = aba.cell(row=1, column=coluna).column_letter
                aba.column_dimensions[letra].width = min(largura + 2, 60)
        buffer = BytesIO()
        livro.save(buffer)
        resposta = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resposta["Content-Disposition"] = f'attachment; filename="{nome}.xlsx"'
        return resposta
    resposta = HttpResponse(content_type="text/csv")
    resposta["Content-Disposition"] = f'attachment; filename="{nome}.csv"'
    resposta.write("\ufeff")
    writer = csv.writer(resposta, delimiter=";")
    writer.writerow(colunas)
    for grupo, valor, entra in linhas:
        writer.writerow([grupo, f"{Decimal(valor):.2f}".replace(".", ","), entra])
    return resposta


def das(request):
    """Base de cálculo do DAS — só receita de locação é tributável (decisão nº 11).

    Sem ?mes=, abre no mês anterior até o dia 20 (prazo do DAS) e no mês
    atual depois disso. Setas de mês no lugar do <input type=month> (Safari).
    """
    hoje = timezone.localdate()
    ano, mes = ano_mes(request, padrao=mes_padrao_do_das(hoje))
    if not ANO_MINIMO <= ano <= hoje.year + 1:
        ano, mes = mes_padrao_do_das(hoje)
    resumo = services.resumo_fiscal(ano, mes)
    formato = request.GET.get("exportar")
    if formato in ("csv", "xlsx"):
        return _planilha_das(
            formato,
            f"das-{ano}-{mes:02d}",
            linhas_base_das(ano, mes),
            linhas_pagamentos_das(ano, mes),
        )
    return render(
        request,
        "financeiro/das.html",
        {
            "resumo": resumo,
            "ano": ano,
            "mes": mes,
            "mes_str": f"{ano}-{mes:02d}",
            **navegacao_de_meses(ano, mes, hoje),
        },
    )


def extrato_cliente(request, cliente_id):
    cliente = get_object_or_404(Cliente, pk=cliente_id)
    hoje = timezone.localdate()
    recebimentos = list(
        services.anotar_lancamento(cliente.recebimentos.all()).prefetch_related(
            "aplicacoes__cobranca", "movimentos_credito"
        )[:50]
    )
    # sobra que virou reforço de caução: achada pela marca no texto (uma consulta)
    reforcos = {
        m.observacoes: m.valor
        for m in MovimentacaoCaucao.objects.filter(
            tipo=MovimentacaoCaucao.Tipo.REFORCO,
            observacoes__in=[services._marca_reforco(r) for r in recebimentos],
        )
    }
    for recebimento in recebimentos:  # lançado hoje: tem o botão "Desfazer"
        recebimento.pode_desfazer = services.pode_desfazer(recebimento)
        entrada = sum(
            (
                m.valor
                for m in recebimento.movimentos_credito.all()
                if m.tipo == MovimentoCredito.Tipo.ENTRADA
            ),
            ZERO,
        )
        if entrada > ZERO:
            recebimento.sobra = (entrada, "crédito")
        elif services._marca_reforco(recebimento) in reforcos:
            recebimento.sobra = (reforcos[services._marca_reforco(recebimento)], "caução")
        else:
            recebimento.sobra = None
    devidas = _cobrancas_devidas(cliente)
    vencido, a_vencer = _vencido_e_a_vencer(devidas, hoje)
    caucao = sum(
        (
            c.saldo
            for c in Caucao.com_saldos_anotados(Caucao.objects.filter(alocacao__cliente=cliente))
        ),
        ZERO,
    )
    return render(
        request,
        "financeiro/extrato_cliente.html",
        {
            "cliente": cliente,
            "cobrancas": Cobranca.com_quitacao_anotada(cliente.cobrancas.order_by("-vencimento"))[
                :100
            ],
            "recebimentos": recebimentos,
            "status_devidos": Cobranca.STATUS_DEVIDOS,  # quem ganha o botão "Recebi"
            "credito": MovimentoCredito.saldo_do_cliente(cliente),
            "vencido": vencido,
            "a_vencer": a_vencer,
            "devido_total": vencido + a_vencer,
            "saldo_caucao": caucao,
            # mensagem única com tudo em aberto (judicial fica de fora)
            "link_whatsapp": whatsapp.link_cliente(
                cliente, _abertas_por_cliente([cliente.pk]).get(cliente.pk, []), hoje
            ),
        },
    )
